from __future__ import annotations

import copy
import csv
import io
import json
import re
import zipfile

import pandas as pd
import pytest
from docx import Document
from pydantic import ValidationError

import streamlit_shared as shared
from app.services.mfi_drafter.synthetic_fixtures import SyntheticSpec, build_profile
from app.services.mfi_drafter.table_projection import (
    MFI_QA_FINDINGS_TABLE_SPEC,
    MFI_REPORT_TABLE_SPECS,
    MFIReportTableColumn,
    MFIReportTableProjectionError,
    MFIReportTableSpec,
    build_mfi_presentation_table,
    build_mfi_qa_presentation_table,
    build_mfi_raw_table_downloads,
    format_mfi_report_value,
    get_mfi_report_table_spec,
)
from app.shared.docx_export import build_docx_bytes_from_report_blocks
from app.shared.report_blocks import ReportBlock


@pytest.fixture(scope="module")
def profile() -> dict:
    return build_profile(
        SyntheticSpec(market_count=20, item_market_ratio=0.5)
    ).model_dump(mode="python")


def _priority_dimension(profile: dict, *, with_items: bool = False) -> str:
    priorities = profile["priority_dimension_names"]
    if not with_items:
        return priorities[0]
    represented = {
        row["values"]["dimension"]
        for row in profile["tables"]["relevant_item_rows"]
    }
    return next(dimension for dimension in priorities if dimension in represented)


def _all_projection_meta(profile: dict) -> list[dict]:
    metadata = [
        build_mfi_presentation_table(
            profile,
            spec_id="mfi.dimension_summary.v1",
            title="Dimension summary",
        ),
        build_mfi_presentation_table(
            profile,
            spec_id="mfi.regional_summary.v1",
            title="Regional summary",
        ),
        build_mfi_presentation_table(
            profile,
            spec_id="mfi.priority_market.v1",
            title="Priority markets",
        ),
    ]
    for dimension in profile["priority_dimension_names"]:
        if dimension != "Food Quality":
            metadata.append(
                build_mfi_presentation_table(
                    profile,
                    spec_id="mfi.official_subsection.v1",
                    title=f"{dimension} subsections",
                    dimension=dimension,
                )
            )
        metadata.extend(
            [
                build_mfi_presentation_table(
                    profile,
                    spec_id="mfi.ranked_driver.v1",
                    title=f"{dimension} drivers",
                    dimension=dimension,
                ),
                build_mfi_presentation_table(
                    profile,
                    spec_id="mfi.relevant_item.v1",
                    title=f"{dimension} items",
                    dimension=dimension,
                ),
            ]
        )
    return metadata


def test_table_specs_are_immutable_unique_and_within_contract() -> None:
    assert len(MFI_REPORT_TABLE_SPECS) == 6
    assert len({spec.spec_id for spec in MFI_REPORT_TABLE_SPECS.values()}) == 6
    assert len(
        {spec.canonical_source_table for spec in MFI_REPORT_TABLE_SPECS.values()}
    ) == 6
    for spec in MFI_REPORT_TABLE_SPECS.values():
        assert 1 <= len(spec.columns) <= 8
        assert len({column.key for column in spec.columns}) == len(spec.columns)
        assert all(column.width_hint > 0 for column in spec.columns)
        assert all(
            column.alignment in {"left", "center", "right"}
            for column in spec.columns
        )
    with pytest.raises(ValidationError):
        get_mfi_report_table_spec("mfi.dimension_summary.v1").maximum_rows = 10


@pytest.mark.parametrize(
    ("kwargs", "message"),
    [
        (
            {
                "spec_id": "duplicate",
                "canonical_source_table": "rows",
                "columns": (
                    MFIReportTableColumn(key="a", label="A"),
                    MFIReportTableColumn(key="a", label="Again"),
                ),
            },
            "unique",
        ),
        (
            {
                "spec_id": "too-wide",
                "canonical_source_table": "rows",
                "columns": tuple(
                    MFIReportTableColumn(key=f"c{i}", label=f"C{i}")
                    for i in range(9)
                ),
            },
            "eight",
        ),
        (
            {
                "spec_id": "bad-cap",
                "canonical_source_table": "rows",
                "columns": (MFIReportTableColumn(key="a", label="A"),),
                "maximum_rows": 0,
            },
            "positive",
        ),
    ],
)
def test_table_spec_validation(kwargs, message) -> None:
    with pytest.raises(ValidationError, match=message):
        MFIReportTableSpec(**kwargs)


def test_column_width_and_text_validation() -> None:
    with pytest.raises(ValidationError, match="positive"):
        MFIReportTableColumn(key="a", label="A", width_hint=0)
    with pytest.raises(ValidationError, match="non-empty"):
        MFIReportTableColumn(key=" ", label="A")


@pytest.mark.parametrize(
    ("value", "format_kind", "expected"),
    [
        (3.14159, "score_statistic", "3.14"),
        (0.31415, "percentage", "31.4%"),
        (0.1049, "percentage_point", "10.5 pp"),
        (3, "integer", "3"),
        (["Price", "Service"], "list", "Price, Service"),
        (True, "boolean", "Yes"),
        (False, "boolean", "No"),
        (None, "text", "—"),
        (
            {
                "available_market_count": 3,
                "total_assessed_market_count": 4,
                "coverage_ratio": 0.75,
            },
            "coverage",
            "3/4 markets (75.0%)",
        ),
    ],
)
def test_shared_formatters(value, format_kind, expected) -> None:
    assert format_mfi_report_value(value, format_kind) == expected


def test_formatters_reject_invalid_values() -> None:
    with pytest.raises(MFIReportTableProjectionError, match="non-integral"):
        format_mfi_report_value(1.5, "integer")
    with pytest.raises(MFIReportTableProjectionError, match="inconsistent"):
        format_mfi_report_value(
            {
                "available_market_count": 1,
                "total_assessed_market_count": 2,
                "coverage_ratio": 0.9,
            },
            "coverage",
        )


def test_dimension_and_regional_projections_are_explicit_and_complete(profile) -> None:
    dimension = build_mfi_presentation_table(
        profile,
        spec_id="mfi.dimension_summary.v1",
        title="Dimension summary",
    )
    regional = build_mfi_presentation_table(
        profile,
        spec_id="mfi.regional_summary.v1",
        title="Regional summary",
    )
    assert dimension["columns"] == [
        "dimension",
        "mean",
        "median",
        "minimum",
        "maximum",
        "iqr",
        "rank",
        "is_priority",
    ]
    assert len(dimension["rows"]) == 9
    assert len(regional["rows"]) == len(profile["tables"]["regional_rows"])
    assert all(
        re.fullmatch(r"\d+/\d+ markets \(\d+\.\d%\)", row["values"]["coverage"])
        for row in regional["rows"]
    )


def test_priority_subsections_and_top_four_drivers(profile) -> None:
    dimension = _priority_dimension(profile)
    subsection = build_mfi_presentation_table(
        profile,
        spec_id="mfi.official_subsection.v1",
        title="Subsections",
        dimension=dimension,
    )
    drivers = build_mfi_presentation_table(
        profile,
        spec_id="mfi.ranked_driver.v1",
        title="Drivers",
        dimension=dimension,
    )
    assert len(subsection["rows"]) <= 2
    assert len(drivers["rows"]) <= 4
    source_by_id = {
        row["row_id"]: row["values"]
        for row in profile["tables"]["driver_rows"]
    }
    assert all(
        source_by_id[row["row_id"]]["role"]
        in {"category_driver", "question_driver"}
        for row in drivers["rows"]
    )
    assert [
        int(row["values"]["weakness_rank"]) for row in drivers["rows"]
    ] == sorted(int(row["values"]["weakness_rank"]) for row in drivers["rows"])


def test_food_quality_uses_question_drivers_and_no_invented_subsections(profile) -> None:
    quality_profile = copy.deepcopy(profile)
    if "Food Quality" not in quality_profile["priority_dimension_names"]:
        quality_profile["priority_dimension_names"].append("Food Quality")
    subsection = build_mfi_presentation_table(
        quality_profile,
        spec_id="mfi.official_subsection.v1",
        title="Quality subsections",
        dimension="Food Quality",
    )
    drivers = build_mfi_presentation_table(
        quality_profile,
        spec_id="mfi.ranked_driver.v1",
        title="Quality questions",
        dimension="Food Quality",
    )
    source_by_id = {
        row["row_id"]: row["values"]
        for row in profile["tables"]["driver_rows"]
    }
    assert subsection["rows"] == []
    assert drivers["rows"]
    assert all(
        source_by_id[row["row_id"]]["role"] == "question_driver"
        for row in drivers["rows"]
    )


def test_relevant_items_only_and_category_contrast_is_derived(profile) -> None:
    dimension = _priority_dimension(profile, with_items=True)
    meta = build_mfi_presentation_table(
        profile,
        spec_id="mfi.relevant_item.v1",
        title="Relevant items",
        dimension=dimension,
    )
    canonical = {
        row["row_id"]: row["values"]
        for row in profile["tables"]["relevant_item_rows"]
        if row["values"]["dimension"] == dimension
    }
    drivers = {
        row["values"]["metric_id"]: row["values"]
        for row in profile["tables"]["driver_rows"]
    }
    assert {row["row_id"] for row in meta["rows"]} == set(canonical)
    for row in meta["rows"]:
        item = canonical[row["row_id"]]
        category = drivers[item["matching_category_metric_id"]]
        expected = (
            float(item["unfavorable_rate"])
            - float(category["unfavorable_rate"])
        )
        assert row["values"]["category_contrast"] == f"{expected * 100:.1f} pp"
        assert item["item_relevant"] is True


def test_price_projections_exclude_items_and_nonqualifying_internal_rows(profile) -> None:
    assert "Price" in profile["priority_dimension_names"]
    drivers = build_mfi_presentation_table(
        profile,
        spec_id="mfi.ranked_driver.v1",
        title="Price drivers",
        dimension="Price",
    )
    source = {
        row["row_id"]: row["values"] for row in profile["tables"]["driver_rows"]
    }
    assert all(source[row["row_id"]]["role"] != "item_driver" for row in drivers["rows"])
    assert all("metric_id" not in row["values"] for row in drivers["rows"])
    items = build_mfi_presentation_table(
        profile,
        spec_id="mfi.relevant_item.v1",
        title="Price items",
        dimension="Price",
    )
    relevant_ids = {
        row["row_id"]
        for row in profile["tables"]["relevant_item_rows"]
        if row["values"]["dimension"] == "Price"
    }
    assert {row["row_id"] for row in items["rows"]} == relevant_ids


def test_priority_market_projection_is_strictly_capped_at_fifteen(profile) -> None:
    meta = build_mfi_presentation_table(
        profile,
        spec_id="mfi.priority_market.v1",
        title="Priority markets",
    )
    assert len(meta["rows"]) == 15
    canonical_order = [
        row["row_id"]
        for row in sorted(
            profile["tables"]["priority_market_rows"],
            key=lambda row: (
                row["values"]["selection_order"],
                row["values"]["market_name"].casefold(),
            ),
        )
    ][:15]
    assert [row["row_id"] for row in meta["rows"]] == canonical_order


def test_projection_does_not_mutate_canonical_profile(profile) -> None:
    before = copy.deepcopy(profile["tables"])
    _all_projection_meta(profile)
    assert profile["tables"] == before


def test_every_visible_analytical_cell_has_valid_ledger_linkage(profile) -> None:
    ledger = profile["metric_ledger"]
    for meta in _all_projection_meta(profile):
        for row in meta["rows"]:
            assert set(row["cell_ledger_metric_ids"]) == set(meta["columns"])
            for column in meta["columns"]:
                links = row["cell_ledger_metric_ids"][column]
                assert links
                assert set(links) <= set(ledger)


def test_projection_fails_for_missing_spec_or_broken_linkage(profile) -> None:
    with pytest.raises(MFIReportTableProjectionError, match="No registered"):
        get_mfi_report_table_spec("mfi.unregistered.v1")
    broken = copy.deepcopy(profile)
    broken_id = broken["tables"]["dimension_rows"][0]["ledger_metric_ids"][0]
    del broken["metric_ledger"][broken_id]
    with pytest.raises(MFIReportTableProjectionError, match="missing ledger"):
        build_mfi_presentation_table(
            broken,
            spec_id="mfi.dimension_summary.v1",
            title="Broken",
        )


def test_nonpriority_evidence_projection_fails_closed(profile) -> None:
    nonpriority = next(
        dimension
        for dimension in (
            row["values"]["dimension"]
            for row in profile["tables"]["dimension_rows"]
        )
        if dimension not in profile["priority_dimension_names"]
    )
    with pytest.raises(MFIReportTableProjectionError, match="priority dimension"):
        build_mfi_presentation_table(
            profile,
            spec_id="mfi.ranked_driver.v1",
            title="Not selected",
            dimension=nonpriority,
        )


def test_qa_findings_projection_is_eight_columns_with_full_raw_metadata() -> None:
    original = {
        "severity": "MEDIUM",
        "source": "deterministic",
        "artifact": "dimension",
        "location": "Price",
        "field": "key_findings",
        "claim_id": "dimension.price.finding.1",
        "code": "scope_mismatch",
        "message": "Review the claim.",
        "attempts": 2,
        "outcome": "partially resolved",
        "disposition": "retained_unverified_for_delivery",
    }
    meta = build_mfi_qa_presentation_table(
        title="QA findings",
        qa_flag_ids=["flag-1"],
        rows=[
            {
                "row_id": "flag-1",
                "values": {
                    "severity": "MEDIUM",
                    "source": "deterministic",
                    "artifact_location": "dimension / Price",
                    "field": "key_findings",
                    "claim_code": "dimension.price.finding.1 / scope_mismatch",
                    "message": "Review the claim.",
                    "attempts_outcome": "2 / partially resolved",
                    "disposition": "retained_unverified_for_delivery",
                },
                "raw_values": original,
            }
        ],
    )
    assert len(meta["columns"]) == 8
    assert meta["column_specs"] == [
        column.model_dump(mode="json") for column in MFI_QA_FINDINGS_TABLE_SPEC.columns
    ]
    assert meta["rows"][0]["raw_values"] == original


def test_raw_downloads_are_complete_full_precision_and_deterministic(profile) -> None:
    first = build_mfi_raw_table_downloads(profile)
    second = build_mfi_raw_table_downloads(profile)
    assert first == second
    assert len(first) == 7
    assert len({item["file_name"] for item in first}) == 7
    bundle = json.loads(first[0]["data"].decode("utf-8"))
    assert bundle == profile["tables"]
    dimension_csv = next(
        item for item in first if item["file_name"] == "mfi_dimension_raw.csv"
    )
    csv_rows = list(
        csv.DictReader(io.StringIO(dimension_csv["data"].decode("utf-8-sig")))
    )
    canonical = profile["tables"]["dimension_rows"][0]
    assert csv_rows[0]["row_id"] == canonical["row_id"]
    assert float(csv_rows[0]["mean"]) == canonical["values"]["mean"]
    assert csv_rows[0]["mean"] != f"{canonical['values']['mean']:.2f}"


def test_streamlit_and_docx_use_the_same_projected_labels_and_values(
    profile, monkeypatch
) -> None:
    meta = build_mfi_presentation_table(
        profile,
        spec_id="mfi.dimension_summary.v1",
        title="Dimension summary",
    )
    captured: list[pd.DataFrame] = []
    monkeypatch.setattr(shared.st, "markdown", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(
        shared.st,
        "dataframe",
        lambda value, **_kwargs: captured.append(value.copy()),
    )
    shared.render_report_blocks(
        [ReportBlock(type="table", meta=meta).model_dump(mode="python")], {}
    )
    assert len(captured) == 1
    labels = [column["label"] for column in meta["column_specs"]]
    assert list(captured[0].columns) == labels
    assert captured[0].iloc[0].tolist() == [
        meta["rows"][0]["values"][column] for column in meta["columns"]
    ]

    docx = build_docx_bytes_from_report_blocks(
        [ReportBlock(type="table", meta=meta)], visualizations={}
    )
    document = Document(io.BytesIO(docx))
    table = document.tables[0]
    assert [cell.text for cell in table.rows[0].cells] == labels
    assert [cell.text for cell in table.rows[1].cells] == [
        meta["rows"][0]["values"][column] for column in meta["columns"]
    ]


def test_docx_projection_xml_has_repeat_headers_fixed_widths_and_non_split_rows(
    profile,
) -> None:
    meta = build_mfi_presentation_table(
        profile,
        spec_id="mfi.dimension_summary.v1",
        title="Dimension summary",
    )
    docx = build_docx_bytes_from_report_blocks(
        [ReportBlock(type="table", meta=meta)], visualizations={}
    )
    with zipfile.ZipFile(io.BytesIO(docx)) as archive:
        xml = archive.read("word/document.xml").decode("utf-8")
    assert "w:tblHeader" in xml
    assert xml.count("w:cantSplit") == len(meta["rows"])
    assert "w:tcW" in xml
    document = Document(io.BytesIO(docx))
    visible = "\n".join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert "metric_id" not in visible
    assert "True" not in visible and "False" not in visible
    assert not re.search(r"\b\d+\.\d{3,}\b", visible)


def test_renderers_refuse_unprojected_mfi_tables(monkeypatch) -> None:
    block = ReportBlock(
        type="table",
        meta={
            "table_kind": "mfi_deterministic",
            "title": "Raw",
            "rows": [{"values": {"mean": 1.23456789}}],
        },
    )
    with pytest.raises(ValueError, match="Unprojected"):
        build_docx_bytes_from_report_blocks([block], visualizations={})
    with pytest.raises(ValueError, match="Unprojected"):
        shared.render_report_blocks([block.model_dump(mode="python")], {})


def test_technical_download_renderer_exposes_json_and_six_csvs(profile, monkeypatch) -> None:
    buttons = []
    monkeypatch.setattr(shared.st, "markdown", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(shared.st, "caption", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(
        shared.st,
        "download_button",
        lambda label, **kwargs: buttons.append((label, kwargs)),
    )
    shared.render_mfi_raw_table_downloads(profile, key_prefix="r6")
    assert len(buttons) == 7
    assert buttons[0][1]["file_name"].endswith(".json")
    assert all(item[1]["file_name"].endswith(".csv") for item in buttons[1:])
    assert len({item[1]["key"] for item in buttons}) == 7
