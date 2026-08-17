from __future__ import annotations

import io

from docx import Document

from app.services.mfi_drafter import graph
from app.services.mfi_drafter.report_inspector import inspect_docx_bytes
from app.services.mfi_drafter.schemas import GenerateMFIReportOutput
from app.shared.docx_export import build_docx_bytes_from_report_blocks
from app.shared.report_blocks import build_mfi_report_blocks
import streamlit_shared as shared


def _flag(
    flag_id: str,
    *,
    severity: str = "medium",
    code: str = "review_required",
    claim_id: str | None = "dimension.price.finding.1",
    artifact_type: str = "dimension",
    artifact_id: str | None = "Price",
    field_name: str | None = "key_findings",
    message: str = "The claim requires review.",
    repairable: bool = True,
) -> dict:
    return {
        "flag_id": flag_id,
        "source": "deterministic",
        "code": code,
        "severity": severity,
        "artifact_type": artifact_type,
        "artifact_id": artifact_id,
        "field_name": field_name,
        "claim_id": claim_id,
        "message": message,
        "recommendation": "Review the affected field.",
        "metric_ids": [],
        "document_ids": [],
        "expected_value": None,
        "actual_value": None,
        "repairable": repairable,
    }


def _claim(text: str, *, substituted: bool = False) -> dict:
    return {
        "claim_id": "dimension.price.finding.1",
        "text": text,
        "claim_kind": "finding",
        "metric_ids": ["assessment.dimension.price.mean"],
        "document_ids": [],
        "scope": "assessment",
        "polarity": "neutral",
        "validation_status": "unverified",
        "validation_flags": [],
        "validation_flag_ids": [],
        "substituted": substituted,
    }


def _result(*, claim: dict | None = None, flags: list[dict] | None = None) -> dict:
    narrative = None
    if claim is not None:
        narrative = {
            "dimension": "Price",
            "is_priority": False,
            "summary": {
                **_claim("Price evidence is summarized below."),
                "claim_id": "dimension.price.summary",
                "validation_status": "verified",
            },
            "key_findings": [claim],
            "subdimension_analysis": [],
            "geographic_patterns": [],
            "data_limitations": [],
            "recommendations": [],
        }
    return {
        "country": "Testland",
        "data_collection_start": "2026-01-01",
        "data_collection_end": "2026-01-31",
        "methodology_version": "databridge-current",
        "score_authority": "synthetic_mock",
        "assessment_profile": {
            "priority_dimension_names": [],
            "priority_market_names": [],
            "markets": [],
            "limitations": [],
            "tables": {},
        },
        "claim_catalog": {
            "assessment.dimension.price.mean": {
                "label": "Price mean",
                "formatted_value": "5.00/10",
                "scope": "assessment",
                "coverage_label": "3/3 assessed markets",
            }
        },
        "context_evidence": [],
        "dimension_narratives": {"Price": narrative} if narrative else {},
        "market_narratives": {},
        "executive_summary_narrative": {},
        "qa_review": {
            "status": "completed_with_warnings" if flags else "passed",
            "correction_attempts": 0,
            "correction_history": [],
            "flags": flags or [],
        },
        "generation_diagnostics": {
            "claim_substitutions": [],
            "unmatched_high_claim_ids": [],
        },
        "visualizations": {},
    }


def _docx_text(docx: bytes) -> str:
    document = Document(io.BytesIO(docx))
    parts = [paragraph.text for paragraph in document.paragraphs]
    parts.extend(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    return "\n".join(parts)


def test_multiple_material_flags_share_one_adjacent_claim_warning() -> None:
    flags = [
        _flag("flag-a", code="scope_mismatch"),
        _flag("flag-b", code="polarity_mismatch"),
    ]
    blocks = build_mfi_report_blocks(_result(claim=_claim("Price scored 5.00/10."), flags=flags))

    claim_index = next(
        index
        for index, block in enumerate(blocks)
        if block.type == "paragraph"
        and (block.meta or {}).get("claim_id") == "dimension.price.finding.1"
    )
    assert blocks[claim_index + 1].type == "evidence_note"
    warning = blocks[claim_index + 2]
    assert warning.type == "claim_warning"
    assert warning.meta["flag_codes"] == ["polarity_mismatch", "scope_mismatch"]
    assert warning.meta["disposition"] == "retained_unverified_for_delivery"
    assert sum(block.type == "claim_warning" for block in blocks) == 1

    qa_table = next(
        block
        for block in blocks
        if block.type == "table" and (block.meta or {}).get("title") == "QA findings"
    )
    assert [row["row_id"] for row in qa_table.meta["rows"]] == ["flag-b", "flag-a"]


def test_high_claim_shows_withdrawal_and_never_exports_rejected_text() -> None:
    rejected = "Use MFI alone to choose cash assistance."
    replacement = "MFI findings do not determine transfer modality."
    flag = _flag(
        "flag-high",
        severity="high",
        code="unsupported_modality_conclusion",
        message="MFI evidence cannot determine transfer modality.",
    )
    result = _result(claim=_claim(replacement, substituted=True), flags=[flag])
    result["generation_diagnostics"]["claim_substitutions"] = [
        {
            "claim_id": "dimension.price.finding.1",
            "rejected_text": rejected,
            "replacement_text": replacement,
            "codes": ["unsupported_modality_conclusion"],
            "flag_ids": ["flag-high"],
            "disposition": "replaced_by_deterministic_fallback",
        }
    ]

    blocks = build_mfi_report_blocks(result)
    warning = next(block for block in blocks if block.type == "claim_warning")
    assert warning.meta["disposition"] == "replaced_by_deterministic_fallback"
    assert "DO NOT USE ORIGINAL" in warning.text

    docx = build_docx_bytes_from_report_blocks(blocks, visualizations={})
    text = _docx_text(docx)
    assert "dimension.price.finding.1" in text
    assert "unsupported_modality_conclusion" in text
    assert rejected not in text
    inspected = inspect_docx_bytes(docx)
    assert inspected.claim_status_marker_count >= 1
    assert inspected.qa_table_row_count == 1


def test_global_red_team_failure_is_listed_without_fake_claim_marker() -> None:
    flag = _flag(
        "system-red-team-execution-error",
        severity="medium",
        code="qa_execution_error",
        claim_id=None,
        artifact_type="global",
        artifact_id=None,
        field_name=None,
        message="The LLM Red-Team review could not be completed.",
        repairable=False,
    )
    blocks = build_mfi_report_blocks(_result(flags=[flag]))

    assert not [block for block in blocks if block.type == "claim_warning"]
    global_notice = next(
        block
        for block in blocks
        if block.type == "qa_warning" and (block.meta or {}).get("global_flag_ids")
    )
    assert global_notice.meta["global_flag_ids"] == [flag["flag_id"]]
    assert "qa_execution_error" in global_notice.text
    qa_table = next(
        block
        for block in blocks
        if block.type == "table" and (block.meta or {}).get("title") == "QA findings"
    )
    assert qa_table.meta["rows"][0]["values"]["claim_id"] == ""
    assert qa_table.meta["rows"][0]["values"]["disposition"] == "global_unresolved"


def test_unmatched_high_claim_id_is_promoted_to_global_delivery_notice() -> None:
    flag = _flag(
        "unmatched-high",
        severity="high",
        code="invalid_metric_id",
        claim_id="claim-that-is-not-rendered",
    )
    blocks = build_mfi_report_blocks(_result(flags=[flag]))

    assert not [block for block in blocks if block.type == "claim_warning"]
    notice = next(
        block
        for block in blocks
        if block.type == "qa_warning" and (block.meta or {}).get("global_flag_ids")
    )
    assert notice.meta["global_flag_ids"] == ["unmatched-high"]
    assert "invalid_metric_id" in notice.text

    finalized = graph.node_finalize_qa(
        {
            "deterministic_flags": [flag],
            "red_team_flags": [],
            "correction_attempts": 0,
            "correction_history": [],
            "dimension_narratives": {},
            "market_narratives": {},
            "executive_summary_narrative": {},
            "context_evidence": [],
        }
    )
    assert any("process-level" in message for message in finalized["warnings"])
    assert not any("claim-level" in message for message in finalized["warnings"])


def test_low_advisory_does_not_create_claim_warning_or_material_table() -> None:
    flag = _flag("low-advisory", severity="low", code="style_advisory")
    result = _result(claim=_claim("Price scored 5.00/10."), flags=[flag])
    result["qa_review"]["status"] = "passed_with_advisories"

    blocks = build_mfi_report_blocks(result)

    assert not [block for block in blocks if block.type == "claim_warning"]
    assert not [
        block
        for block in blocks
        if block.type == "table" and (block.meta or {}).get("title") == "QA findings"
    ]


def test_streamlit_preview_renders_claim_warning_and_qa_table(monkeypatch) -> None:
    flag = _flag("preview-medium", code="scope_mismatch")
    blocks = build_mfi_report_blocks(
        _result(claim=_claim("Price scored 5.00/10."), flags=[flag])
    )
    events: list[tuple[str, object]] = []
    monkeypatch.setattr(shared.st, "warning", lambda value: events.append(("warning", value)))
    monkeypatch.setattr(shared.st, "markdown", lambda value: events.append(("markdown", value)))
    monkeypatch.setattr(
        shared.st,
        "dataframe",
        lambda value, **_kwargs: events.append(("dataframe", value)),
    )

    shared.render_report_blocks(
        [
            block.model_dump()
            for block in blocks
            if block.type in {"claim_warning", "table"}
            and (
                block.type == "claim_warning"
                or (block.meta or {}).get("title") == "QA findings"
            )
        ],
        {},
    )

    assert any(
        kind == "warning" and "dimension.price.finding.1" in str(value)
        for kind, value in events
    )
    assert any(kind == "dataframe" for kind, _value in events)


def test_failed_targeted_repair_is_counted_and_retained_in_final_review() -> None:
    flag = _flag("flag-failed", severity="high", code="invalid_metric_id")
    prepared = graph.node_prepare_correction(
        {
            "deterministic_flags": [flag],
            "red_team_flags": [],
            "correction_attempts": 0,
            "correction_history": [],
            "contextual_documents": [],
            "context_evidence": [],
        }
    )
    history = graph._record_correction_execution(
        prepared["correction_history"],
        attempt_number=1,
        targets=prepared["correction_targets"],
        outcome="llm_or_schema_failed",
    )
    final = graph.node_finalize_qa(
        {
            "deterministic_flags": [flag],
            "red_team_flags": [],
            "correction_attempts": 1,
            "correction_history": history,
            "dimension_narratives": {
                "Price": {"key_findings": [_claim("Invalid drafted claim.")]}
            },
            "market_narratives": {},
            "executive_summary_narrative": {},
            "context_evidence": [],
        }
    )

    record = final["qa_review"]["correction_history"][0]
    assert record["claim_id"] == "dimension.price.finding.1"
    assert record["execution_outcome"] == "llm_or_schema_failed"
    assert record["validation_outcome"] == "unresolved"
    assert final["qa_review"]["correction_attempts"] == 1

    result = _result(
        claim=final["dimension_narratives"]["Price"]["key_findings"][0],
        flags=final["qa_review"]["flags"],
    )
    result["qa_review"] = final["qa_review"]
    result["generation_diagnostics"] = final["generation_diagnostics"]
    warning = next(
        block
        for block in build_mfi_report_blocks(result)
        if block.type == "claim_warning"
    )
    assert warning.meta["repair_attempted"] is True
    assert warning.meta["attempt_count"] == 1
    assert "llm_or_schema_failed / unresolved" in warning.text


def test_high_context_statement_is_withdrawn_but_kept_in_technical_diagnostics() -> None:
    rejected = "The price shock caused market failure."
    flag = _flag(
        "context-causal",
        severity="high",
        code="unsupported_causal_claim",
        claim_id=None,
        artifact_type="context",
        artifact_id="context-1",
        field_name="text",
    )
    final = graph.node_finalize_qa(
        {
            "deterministic_flags": [flag],
            "red_team_flags": [],
            "correction_attempts": 0,
            "correction_history": [],
            "dimension_narratives": {},
            "market_narratives": {},
            "executive_summary_narrative": {},
            "context_evidence": [
                {
                    "statement_id": "context-1",
                    "text": rejected,
                    "classification": "potentially_explanatory",
                    "document_ids": ["doc-1"],
                    "validation_status": "unverified",
                    "validation_flags": ["unsupported_causal_claim"],
                }
            ],
        }
    )

    statement = final["context_evidence"][0]
    assert statement["substituted"] is True
    assert rejected not in statement["text"]
    substitution = final["generation_diagnostics"]["claim_substitutions"][0]
    assert substitution["rejected_text"] == rejected


def test_public_schema_preserves_r4_audit_fields() -> None:
    definitions = GenerateMFIReportOutput.model_json_schema()["$defs"]

    assert "correction_history" in definitions["MFIQAReview"]["properties"]
    diagnostics = definitions["MFIGenerationDiagnostics"]["properties"]
    assert "claim_substitutions" in diagnostics
    assert "unmatched_high_claim_ids" in diagnostics
