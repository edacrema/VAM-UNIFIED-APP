from __future__ import annotations

import copy
import io
import json
import random
from collections import Counter
from pathlib import Path

import pytest
from docx import Document

from app.services.mfi_drafter import graph
from app.services.mfi_drafter import narrative as narrative_module
from app.services.mfi_drafter.analysis import build_assessment_profile
from app.services.mfi_drafter.errors import MFIGenerationBlockedError
from app.services.mfi_drafter.narrative import (
    build_claim_catalog,
    build_qa_review,
    compact_catalog,
    fallback_dimension_narrative,
    fallback_executive_narrative,
    fallback_market_narrative,
    normalize_red_team_flags,
    validate_structured_narratives,
)
from app.services.mfi_drafter.router import _build_mfi_output
from app.streamlit_backend.dispatcher import _build_mfi_report_output
from app.shared.docx_export import build_docx_bytes_from_report_blocks
from app.shared.llm_observability import LLMCallError
from app.shared.report_blocks import (
    ReportBlock,
    build_mfi_report_blocks,
    resolve_mfi_report_blocks,
)


@pytest.fixture(scope="module")
def phase3_bundle() -> dict:
    random.seed(143)
    data = graph.generate_mock_mfi_data(
        "Testland",
        ["Alpha", "Bravo", "Charlie", "Delta"],
        "2026-01-01",
        "2026-01-31",
    )
    profile = build_assessment_profile(
        data["markets_data"],
        data["metric_summaries"],
        data,
    ).model_dump()
    catalog = build_claim_catalog(profile)
    dimensions = {
        item["dimension"]: fallback_dimension_narrative(
            item,
            assessment_profile=profile,
        )
        for item in profile["dimensions"]
    }
    markets = {
        item["market_name"]: fallback_market_narrative(item)
        for item in profile["markets"]
        if item["is_priority_market"]
    }
    executive = fallback_executive_narrative(profile)
    validation, dimensions, markets, executive, context, flags = (
        validate_structured_narratives(
            context_evidence=[],
            dimension_narratives=dimensions,
            market_narratives=markets,
            executive_narrative=executive,
            claim_catalog=catalog,
            assessment_profile=profile,
            documents=[],
        )
    )
    assert validation["status"] == "passed"
    result = {
        **data,
        "run_id": "phase3-test",
        "country": "Testland",
        "data_collection_start": "2026-01-01",
        "data_collection_end": "2026-01-31",
        "narrative_schema_version": "2.0",
        "release_control": {
            "analysis_version": "2",
            "enabled": True,
            "configuration_status": "configured",
            "service_name": "mfi-drafter",
            "deployment_revision": "phase4-test",
        },
        "generation_diagnostics": {
            "dimensions": {"llm": [], "fallback": sorted(dimensions)},
            "markets": {"llm": [], "fallback": sorted(markets)},
            "context_extraction_mode": "not_applicable",
            "executive_summary_mode": "fallback",
            "red_team_status": "not_started",
            "correction_attempts": 0,
            "unresolved_high_count": 0,
            "unresolved_medium_count": 0,
            "unresolved_low_count": 0,
            "retrievers": {},
        },
        "assessment_profile": profile,
        "mean_mfi_across_assessed_markets": profile[
            "mean_mfi_across_assessed_markets"
        ],
        "claim_catalog": catalog,
        "context_evidence": context,
        "dimension_narratives": dimensions,
        "market_narratives": markets,
        "executive_summary_narrative": executive,
        "claim_validation": validation,
        "qa_review": build_qa_review(
            flags["flags"], [], correction_attempts=0
        ),
        "visualizations": {},
        "document_references": [],
        "llm_calls": 0,
        "correction_attempts": 0,
    }
    return {
        "data": data,
        "profile": profile,
        "catalog": catalog,
        "dimensions": dimensions,
        "markets": markets,
        "executive": executive,
        "result": result,
    }


def _validate(bundle: dict, **updates):
    return validate_structured_narratives(
        context_evidence=updates.get("context_evidence", []),
        dimension_narratives=updates.get(
            "dimension_narratives", copy.deepcopy(bundle["dimensions"])
        ),
        market_narratives=updates.get(
            "market_narratives", copy.deepcopy(bundle["markets"])
        ),
        executive_narrative=updates.get(
            "executive_narrative", copy.deepcopy(bundle["executive"])
        ),
        claim_catalog=bundle["catalog"],
        assessment_profile=bundle["profile"],
        documents=updates.get("documents", []),
    )


def _codes(validation_tuple) -> Counter:
    return Counter(flag["code"] for flag in validation_tuple[0]["flags"])


def test_catalog_uses_closed_phase3_formatting_policy(phase3_bundle):
    catalog = phase3_bundle["catalog"]
    assert catalog["assessment.mfi.mean"]["formatted_value"].endswith("/10")
    assert catalog["assessment.mfi.mean"]["formatted_value"].split("/")[0].count(
        "."
    ) == 1
    assert (
        len(catalog["assessment.mfi.mean"]["formatted_value"].split(".")[-1].split("/")[0])
        == 2
    )
    assert catalog["assessment.mfi.coverage"]["formatted_value"].endswith("%")
    assert len(
        catalog["assessment.mfi.coverage"]["formatted_value"]
        .removesuffix("%")
        .split(".")[-1]
    ) == 1
    assert catalog["assessment.mfi.denominator"]["formatted_value"].isdigit()


def test_priority_markets_have_market_scoped_explanatory_ledger(phase3_bundle):
    ledger = phase3_bundle["profile"]["metric_ledger"]
    for market in phase3_bundle["profile"]["markets"]:
        if not market["is_priority_market"]:
            continue
        weak_dimensions = {
            item["dimension"] for item in market["weak_dimensions"]
        }
        scoped = [
            entry
            for entry in ledger.values()
            if entry.get("market_name") == market["market_name"]
            and entry.get("dimension") in weak_dimensions
            and entry.get("statistic")
            in {
                "market_explanatory_raw_value",
                "market_explanatory_normalized_value",
                "derived_market_unfavorable_rate",
            }
        ]
        assert scoped
        assert all(entry["evidence_scope"] for entry in scoped)


def test_validator_accepts_authorized_display_rounding(phase3_bundle):
    executive = copy.deepcopy(phase3_bundle["executive"])
    entry = phase3_bundle["catalog"]["assessment.mfi.mean"]
    executive["motivation"]["text"] = (
        f"The assessed-market mean was {entry['formatted_value']}."
    )
    executive["motivation"]["metric_ids"] = ["assessment.mfi.mean"]
    validation = _validate(phase3_bundle, executive_narrative=executive)[0]
    assert validation["status"] == "passed"


@pytest.mark.parametrize(
    ("mutation", "expected_code"),
    [
        ("numeric", "numeric_value_mismatch"),
        ("uncited", "uncited_numeric_value"),
        ("invalid_id", "invalid_metric_id"),
        ("unit", "unit_mismatch"),
        ("scope", "scope_mismatch"),
        ("rank", "numeric_value_mismatch"),
        ("terminology", "unsupported_national_terminology"),
        ("modality", "unsupported_modality_conclusion"),
    ],
)
def test_validator_rejects_invalid_claim_contracts(
    phase3_bundle, mutation, expected_code
):
    executive = copy.deepcopy(phase3_bundle["executive"])
    claim = executive["motivation"]
    if mutation == "numeric":
        claim["text"] = "The assessed-market mean was 99.99/10."
        claim["metric_ids"] = ["assessment.mfi.mean"]
    elif mutation == "uncited":
        claim["text"] = "The assessed-market mean was 5.55/10."
        claim["metric_ids"] = []
    elif mutation == "invalid_id":
        claim["metric_ids"] = ["missing.metric"]
    elif mutation == "unit":
        claim["text"] = "The assessed-market mean was 50.0%."
        claim["metric_ids"] = ["assessment.mfi.mean"]
    elif mutation == "scope":
        claim["scope"] = "region"
        claim["metric_ids"] = ["assessment.mfi.mean"]
    elif mutation == "rank":
        dimension = phase3_bundle["profile"]["priority_dimension_names"][0]
        narrative = copy.deepcopy(phase3_bundle["dimensions"])
        dim_claim = narrative[dimension]["summary"]
        dimension_profile = next(
            item
            for item in phase3_bundle["profile"]["dimensions"]
            if item["dimension"] == dimension
        )
        rank_id = next(
            metric_id
            for metric_id in dimension_profile["ledger_metric_ids"]
            if metric_id.endswith(".rank")
        )
        dim_claim["text"] = "This dimension has rank 99."
        dim_claim["metric_ids"] = [rank_id]
        assert expected_code in _codes(
            _validate(phase3_bundle, dimension_narratives=narrative)
        )
        return
    elif mutation == "terminology":
        claim["text"] = "The national score summarizes the assessment."
    elif mutation == "modality":
        claim["text"] = "CBT is feasible based on this assessment."
    assert expected_code in _codes(
        _validate(phase3_bundle, executive_narrative=executive)
    )


def test_validator_accepts_dates_only_from_cited_document(phase3_bundle):
    executive = copy.deepcopy(phase3_bundle["executive"])
    executive["motivation"]["text"] = "A contextual source was published in 2026."
    executive["motivation"]["metric_ids"] = []
    executive["motivation"]["document_ids"] = ["doc-2026"]
    documents = [
        {
            "doc_id": "doc-2026",
            "date": "2026-02-14",
            "title": "Market update",
            "content": "Published update.",
        }
    ]
    validation = _validate(
        phase3_bundle,
        executive_narrative=executive,
        documents=documents,
    )[0]
    assert "numeric_value_mismatch" not in {
        flag["code"] for flag in validation["flags"]
    }


def test_context_contract_forbids_unknown_sources_and_causality(phase3_bundle):
    documents = [
        {
            "doc_id": "doc-1",
            "date": "2026-02-14",
            "title": "Market update",
            "content": "Context.",
        }
    ]
    context = [
        {
            "statement_id": "context-1",
            "text": "The event caused the observed MFI result.",
            "classification": "potentially_explanatory",
            "document_ids": ["doc-1"],
            "validation_status": "pending",
            "validation_flags": [],
        },
        {
            "statement_id": "context-2",
            "text": "A second contextual statement.",
            "classification": "corroborating",
            "document_ids": ["unknown-doc"],
            "validation_status": "pending",
            "validation_flags": [],
        },
    ]
    codes = _codes(
        _validate(
            phase3_bundle,
            context_evidence=context,
            documents=documents,
        )
    )
    assert codes["unsupported_causal_claim"]
    assert codes["invalid_document_id"]


def test_validator_checks_polarity_and_item_eligibility(phase3_bundle):
    dimensions = copy.deepcopy(phase3_bundle["dimensions"])
    priority = phase3_bundle["profile"]["priority_dimension_names"][0]
    narrative = dimensions[priority]
    candidate = next(
        (
            entry
            for entry in phase3_bundle["catalog"].values()
            if entry.get("dimension") == priority
            and entry.get("orientation") == "higher_is_worse"
            and entry.get("unit") == "proportion"
        ),
        None,
    )
    assert candidate is not None
    narrative["key_findings"][0]["metric_ids"] = [candidate["metric_id"]]
    narrative["key_findings"][0]["polarity"] = "favorable"
    assert "polarity_mismatch" in _codes(
        _validate(phase3_bundle, dimension_narratives=dimensions)
    )

    relevant = {
        metric["metric_id"]
        for dimension in phase3_bundle["profile"]["dimensions"]
        for metric in dimension["drivers"]
        if metric["item_relevant"]
    }
    ineligible = next(
        entry
        for entry in phase3_bundle["catalog"].values()
        if any(".item." in source for source in entry["source_metric_ids"])
        and not set(entry["source_metric_ids"]) & relevant
    )
    dimensions = copy.deepcopy(phase3_bundle["dimensions"])
    target = dimensions[ineligible["dimension"]]["key_findings"][0]
    target["metric_ids"] = [ineligible["metric_id"]]
    assert "insufficient_item_coverage" in _codes(
        _validate(phase3_bundle, dimension_narratives=dimensions)
    )


def test_deterministic_validator_makes_no_llm_call(monkeypatch, phase3_bundle):
    monkeypatch.setattr(
        graph,
        "get_model",
        lambda: (_ for _ in ()).throw(AssertionError("LLM requested")),
    )
    update = graph.node_deterministic_claim_validator(
        {
            "context_evidence": [],
            "dimension_narratives": phase3_bundle["dimensions"],
            "market_narratives": phase3_bundle["markets"],
            "executive_summary_narrative": phase3_bundle["executive"],
            "claim_catalog": phase3_bundle["catalog"],
            "assessment_profile": phase3_bundle["profile"],
            "contextual_documents": [],
        }
    )
    assert update["claim_validation"]["status"] == "passed"


def test_red_team_v5_package_is_complete_deterministic_and_compact(
    phase3_bundle,
) -> None:
    state = {
        **phase3_bundle["result"],
        "deterministic_flags": [],
        "contextual_documents": [],
    }
    package = graph._build_red_team_review_package(state)
    repeated = graph._build_red_team_review_package(copy.deepcopy(state))
    assert package == repeated
    assert package["contract_version"] == "mfi-red-team-input-v5"

    claim_ids = [item["id"] for item in package["claims"]]
    assert claim_ids
    assert len(claim_ids) == len(set(claim_ids))
    cited_ids = {
        str(metric_id)
        for claim in package["claims"]
        for metric_id in claim.get("m", [])
    }
    evidence_ids = set(package["evidence_by_metric_id"])
    assert evidence_ids == cited_ids & set(phase3_bundle["catalog"])

    compact = json.dumps(package, sort_keys=True, separators=(",", ":"))
    legacy = json.dumps(
        {
            "dimension_narratives": state["dimension_narratives"],
            "market_narratives": state["market_narratives"],
            "executive_summary": state["executive_summary_narrative"],
            "cited_catalog": compact_catalog(
                phase3_bundle["catalog"], sorted(cited_ids)
            ),
            "context_evidence": state["context_evidence"],
            "limitations": phase3_bundle["profile"]["limitations"],
            "deterministic_flags": [],
        },
        sort_keys=True,
    )
    assert len(compact) < len(legacy) * 0.85


def test_red_team_flag_identity_is_application_owned_and_collision_safe() -> None:
    supplied = {
        "flag_id": "model-controlled-id",
        "code": "semantic_scope_error",
        "severity": "medium",
        "artifact_type": "dimension",
        "artifact_id": "Price",
        "field_name": "key_findings",
        "claim_id": "dimension-price-key-finding-1",
        "message": "The scope requires qualification.",
        "recommendation": "Qualify the claim scope.",
        "metric_ids": ["metric-price"],
        "document_ids": [],
        "repairable": True,
    }
    first = normalize_red_team_flags({"flags": [supplied]})[0]
    repeated = normalize_red_team_flags({"flags": [copy.deepcopy(supplied)]})[0]
    duplicated = normalize_red_team_flags({"flags": [supplied, supplied]})

    assert first["flag_id"] == repeated["flag_id"]
    assert first["flag_id"].startswith("red-team-semantic_scope_error-")
    assert first["flag_id"] != supplied["flag_id"]
    assert duplicated[0]["flag_id"] == duplicated[1]["flag_id"]


def test_deterministic_flag_identity_includes_evidence_context() -> None:
    base = {
        "code": "subsection_mismatch",
        "severity": "medium",
        "artifact_type": "dimension",
        "artifact_id": "Price",
        "field_name": "subdimension_analysis",
        "claim_id": "dimension-price-subdimension-1",
        "message": "The cited subsection is inconsistent with stored evidence.",
    }
    first = narrative_module._flag(**base, metric_ids=["price-increase"])
    second = narrative_module._flag(**base, metric_ids=["price-stability"])

    assert first["flag_id"] != second["flag_id"]


def test_red_team_v6_uses_persisted_batches_and_operation_specific_timeout(
    monkeypatch,
    phase3_bundle,
) -> None:
    observed = {}

    class Model:
        def bind(self, **kwargs):
            observed["bind_kwargs"] = kwargs
            return self

        def invoke(self, messages):
            observed["prompt"] = messages[0].content
            return type("Response", (), {"content": '{"flags": []}'})()

    def model_factory(**kwargs):
        observed["model_kwargs"] = kwargs
        return Model()

    monkeypatch.setattr(graph, "get_model", model_factory)
    state = {
        **phase3_bundle["result"],
        "deterministic_flags": [],
        "contextual_documents": [],
        "llm_diagnostics": {},
    }
    prepared = graph.node_red_team(state)
    assert prepared["red_team_queue"]
    update = graph.node_process_red_team_batch({**state, **prepared})
    assert observed["model_kwargs"]["timeout_seconds"] == 600.0
    assert "REVIEW_BATCH" in observed["prompt"]
    assert observed["bind_kwargs"]["response_mime_type"] == "application/json"
    assert observed["bind_kwargs"]["response_schema"]["required"] == ["flags"]
    call = update["llm_diagnostics"]["calls"][-1]
    assert call["operation"] == "mfi.red_team_review.v6"
    assert call["configured_timeout_seconds"] == 600.0
    assert call["configured_max_retries"] == 2
    assert update["generation_diagnostics"]["red_team_status"] == "in_progress"
    assert update["generation_diagnostics"]["red_team_structured_output"] is True
    assert update["generation_diagnostics"]["red_team_package_within_target"] is True


def test_red_team_v6_normalizes_one_malformed_json_response_in_memory(
    monkeypatch,
    phase3_bundle,
) -> None:
    responses = [
        '{"flags": []}\n{"flags": []}',
        '{"flags": []}',
    ]
    observed = {"operations": []}

    class Model:
        def bind(self, **_kwargs):
            return self

        def invoke(self, _messages):
            return type("Response", (), {"content": responses.pop(0)})()

    monkeypatch.setattr(graph, "get_model", lambda **_kwargs: Model())
    state = {
        **phase3_bundle["result"],
        "run_id": "red-team-format-repair",
        "deterministic_flags": [],
        "contextual_documents": [],
        "llm_diagnostics": {},
    }
    prepared = graph.node_red_team(state)
    update = graph.node_process_red_team_batch({**state, **prepared})

    calls = update["llm_diagnostics"]["calls"]
    assert [call["operation"] for call in calls] == [
        "mfi.red_team_review.v6",
        "mfi.red_team_review.v6.json_normalization.v1",
    ]
    assert calls[0]["status"] == "recovered"
    assert calls[0]["json_root_value_count"] == 2
    assert calls[0]["json_trailing_character_count"] > 0
    assert calls[1]["status"] == "succeeded"
    assert update["llm_diagnostics"]["status"] == "completed"
    assert update["llm_diagnostics"]["recovered_calls"] == 1
    assert update["llm_diagnostics"]["failed_calls"] == 0
    diagnostics = update["generation_diagnostics"]
    assert diagnostics["red_team_format_repair_attempted"] is True
    assert diagnostics["red_team_format_repair_status"] == "completed"
    assert diagnostics["red_team_initial_call_id"] == calls[0]["call_id"]
    assert diagnostics["red_team_format_repair_call_id"] == calls[1]["call_id"]
    assert update["llm_calls"] == state.get("llm_calls", 0) + 2


def test_red_team_v6_fails_after_one_unsuccessful_format_normalization(
    monkeypatch,
    phase3_bundle,
) -> None:
    responses = [
        '{"flags": []}{"flags": []}',
        "still not json",
    ]

    class Model:
        def bind(self, **_kwargs):
            return self

        def invoke(self, _messages):
            return type("Response", (), {"content": responses.pop(0)})()

    monkeypatch.setattr(graph, "get_model", lambda **_kwargs: Model())
    state = {
        **phase3_bundle["result"],
        "run_id": "red-team-format-repair-failed",
        "deterministic_flags": [],
        "contextual_documents": [],
        "llm_diagnostics": {},
    }
    with pytest.raises(LLMCallError) as caught:
        prepared = graph.node_red_team(state)
        graph.node_process_red_team_batch({**state, **prepared})

    assert caught.value.failure_code == "llm_invalid_json"
    assert caught.value.operation == "mfi.red_team_review.v6.json_normalization.v1"
    assert caught.value.batch_id


def test_red_team_v6_does_not_normalize_semantically_incomplete_json(
    monkeypatch,
    phase3_bundle,
) -> None:
    model_calls = 0

    class Model:
        def bind(self, **_kwargs):
            return self

        def invoke(self, _messages):
            nonlocal model_calls
            model_calls += 1
            return type("Response", (), {"content": '{"flags": [{"code": "x"}]}'})()

    monkeypatch.setattr(graph, "get_model", lambda **_kwargs: Model())
    state = {
        **phase3_bundle["result"],
        "run_id": "red-team-contract-failed",
        "deterministic_flags": [],
        "contextual_documents": [],
        "llm_diagnostics": {},
    }
    with pytest.raises(LLMCallError) as caught:
        prepared = graph.node_red_team(state)
        graph.node_process_red_team_batch({**state, **prepared})

    assert caught.value.failure_code == "llm_response_contract_error"
    assert caught.value.operation == "mfi.red_team_review.v6"
    assert model_calls == 1


def test_red_team_v5_gaza_scale_is_partitioned_inside_character_target() -> None:
    catalog = {}
    dimensions = {}
    dimension_profiles = []
    markets = {}
    market_profiles = []
    metric_index = 0

    def claim(claim_id: str) -> dict:
        nonlocal metric_index
        metric_index += 1
        metric_id = f"metric-{metric_index:03d}"
        catalog[metric_id] = {
            "metric_id": metric_id,
            "label": f"Evidence {metric_index}",
            "formatted_value": f"{metric_index % 10}.00",
            "unit": "score_0_10",
            "orientation": "higher_is_better",
            "scope": "assessment",
            "coverage_label": "27/27 assessed markets",
            "representation_kind": "fixed_metric",
            "representation_complete": True,
            "representation_required": False,
        }
        return {
            "claim_id": claim_id,
            "text": f"Analytical statement {claim_id} supported by supplied evidence.",
            "metric_ids": [metric_id],
            "document_ids": [],
            "scope": "assessment",
            "polarity": "neutral",
        }

    for dimension_position in range(1, 10):
        name = f"Dimension {dimension_position}"
        dimensions[name] = {
            "key_findings": [
                claim(f"dimension-{dimension_position}-finding-{position}")
                for position in range(1, 5)
            ]
        }
        dimension_profiles.append(
            {
                "dimension": name,
                "is_priority": dimension_position <= 4,
                "profile_rank": dimension_position,
            }
        )
    for market_position in range(1, 16):
        name = f"Market {market_position:02d}"
        markets[name] = {
            "priority_issues": [
                claim(f"market-{market_position}-issue-{position}")
                for position in range(1, 6)
            ]
        }
        market_profiles.append(
            {
                "market_name": name,
                "is_priority_market": True,
                "score_rank": market_position,
                "selection_order": market_position,
                "weak_dimensions": [{"dimension": "Dimension 1"}],
            }
        )
    executive = {
        "key_findings": [
            claim(f"executive-finding-{position}") for position in range(1, 3)
        ]
    }
    state = {
        "assessment_profile": {
            "dimensions": dimension_profiles,
            "markets": market_profiles,
            "priority_dimension_names": [f"Dimension {position}" for position in range(1, 5)],
            "priority_market_names": [f"Market {position:02d}" for position in range(1, 16)],
            "limitations": [],
        },
        "dimension_narratives": dimensions,
        "market_narratives": markets,
        "executive_summary_narrative": executive,
        "claim_catalog": catalog,
        "context_evidence": [],
        "contextual_documents": [],
        "deterministic_flags": [
            {
                "code": f"deterministic-{position:02d}",
                "severity": "medium",
                "artifact_type": "dimension",
                "artifact_id": "Dimension 1",
                "field_name": "key_findings",
                "claim_id": "dimension-1-finding-1",
                "message": "Existing deterministic finding already handled.",
            }
            for position in range(1, 37)
        ],
    }

    package = graph._build_red_team_review_package(state)
    assert len(package["claims"]) == 113
    batches = graph.build_red_team_batches(package)
    assert batches
    assert all(
        batch.character_count <= graph.MFI_RED_TEAM_BATCH_TARGET_CHARACTERS
        for batch in batches
    )
    local_claim_ids = [
        claim_id
        for batch in batches
        if batch.batch_kind == "local"
        for claim_id in batch.claim_ids
    ]
    assert len(local_claim_ids) == 113
    assert len(local_claim_ids) == len(set(local_claim_ids))


def test_schema_failure_interrupts_enabled_llm_stage(
    monkeypatch, phase3_bundle
):
    class InvalidModel:
        def invoke(self, _messages):
            return type("Response", (), {"content": "not valid JSON"})()

    monkeypatch.setattr(graph, "get_model", lambda: InvalidModel())
    state = {
        "run_id": "phase3-invalid-schema",
        "assessment_profile": phase3_bundle["profile"],
        "claim_catalog": phase3_bundle["catalog"],
        "dimension_narratives": {},
        "market_narratives": {},
        "executive_summary_narrative": {},
        "correction_targets": [],
        "llm_calls": 0,
    }
    with pytest.raises(LLMCallError) as caught:
        graph.node_dimension_drafter(state)
    assert caught.value.failure_code == "llm_invalid_json"
    assert caught.value.node == "dimension_drafter"


def test_field_correction_prompt_is_closed_catalog_and_patch_is_targeted(
    monkeypatch, phase3_bundle
):
    dimension = phase3_bundle["profile"]["priority_dimension_names"][0]
    profile = next(
        item
        for item in phase3_bundle["profile"]["dimensions"]
        if item["dimension"] == dimension
    )
    drafted = copy.deepcopy(phase3_bundle["dimensions"][dimension])
    replacement = copy.deepcopy(drafted["recommendations"][0])
    replacement["text"] = "Use the cited evidence for review."

    class Model:
        def __init__(self):
            self.prompts = []

        def invoke(self, messages):
            self.prompts.append(messages[0].content)
            return type(
                "Response",
                (),
                {"content": json.dumps({"replacement": [replacement]})},
            )()

    model = Model()
    monkeypatch.setattr(graph, "get_model", lambda: model)
    previous = copy.deepcopy(phase3_bundle["dimensions"][dimension])
    flag = {
        "flag_id": "flag-1",
        "source": "deterministic",
        "code": "test_repair",
        "severity": "medium",
        "artifact_type": "dimension",
        "artifact_id": dimension,
        "field_name": "recommendations",
        "claim_id": previous["recommendations"][0]["claim_id"],
        "message": "Rewrite the recommendation.",
        "repairable": True,
    }
    task = {
        "task_id": "task-1",
        "attempt_number": 1,
        "artifact_type": "dimension",
        "artifact_id": dimension,
        "field_name": "recommendations",
        "claim_ids": [previous["recommendations"][0]["claim_id"]],
        "flag_ids": ["flag-1"],
        "flag_codes": ["test_repair"],
    }
    update = graph.node_process_correction_task(
        {
            "assessment_profile": phase3_bundle["profile"],
            "claim_catalog": phase3_bundle["catalog"],
            "dimension_narratives": phase3_bundle["dimensions"],
            "market_narratives": phase3_bundle["markets"],
            "executive_summary_narrative": phase3_bundle["executive"],
            "context_evidence": [],
            "contextual_documents": [],
            "deterministic_flags": [flag],
            "red_team_flags": [],
            "correction_queue": [task],
            "correction_history": [],
            "generation_diagnostics": {},
            "llm_calls": 0,
        }
    )
    assert len(model.prompts) == 1
    assert "AUTHORIZED_CLAIM_CATALOG" in model.prompts[0]
    assert "PATCH_CONTRACT" in model.prompts[0]
    assert '"maximum_items":' in model.prompts[0]
    assert "metric_id" in model.prompts[0]
    assert "sub_scores" not in model.prompts[0]
    assert "validation_status" not in model.prompts[0]
    assert "validation_flags" not in model.prompts[0]
    assert "validation_flag_ids" not in model.prompts[0]
    assert "substituted" not in model.prompts[0]
    assert (
        update["dimension_narratives"][dimension]["key_findings"]
        == previous["key_findings"]
    )
    assert (
        update["dimension_narratives"][dimension]["recommendations"][0]["text"]
        == "Use the cited evidence for review."
    )
    assert (
        update["generation_diagnostics"][
            "ignored_correction_metadata_field_count"
        ]
        == 5
    )
    assert update["llm_diagnostics"]["calls"][-1]["operation"].endswith(
        "_field_correction.v3"
    )


def test_correction_attempt_limit_blocks_unresolved_material_qa():
    flag = {
        "flag_id": "material-1",
        "source": "deterministic",
        "code": "bad_claim",
        "severity": "high",
        "artifact_type": "executive_summary",
        "artifact_id": "executive_summary",
        "field_name": "motivation",
        "claim_id": "executive.motivation.1",
        "message": "Invalid claim.",
        "recommendation": "Repair it.",
        "metric_ids": [],
        "document_ids": [],
        "expected_value": None,
        "actual_value": None,
        "repairable": True,
    }
    assert (
        graph.route_after_deterministic_validation(
            {
                "deterministic_flags": [flag],
                "red_team_flags": [],
                "correction_attempts": 2,
            }
        )
        == "correct"
    )
    with pytest.raises(MFIGenerationBlockedError) as caught:
        graph.route_after_deterministic_validation(
            {
                "deterministic_flags": [flag],
                "red_team_flags": [],
                "correction_attempts": 3,
            }
        )
    assert caught.value.code == "mfi_narrative_qa_unresolved"


def test_graph_runs_simplified_review_and_single_correction_path():
    edges = {
        (edge.source, edge.target)
        for edge in graph.build_graph().get_graph().edges
    }
    assert (
        "executive_summary_drafter",
        "deterministic_claim_validator",
    ) in edges
    assert ("deterministic_claim_validator", "semantic_review") in edges
    assert ("semantic_review", "consolidated_correction") in edges
    assert ("semantic_review", "finalize_qa") in edges
    assert ("consolidated_correction", "post_correction_validator") in edges
    assert (
        "post_correction_validator",
        "corrected_claim_verification",
    ) in edges
    assert ("corrected_claim_verification", "finalize_qa") in edges
    node_ids = set(graph.build_graph().get_graph().nodes)
    assert not {
        "red_team_batch",
        "red_team_finalize",
        "targeted_correction",
        "correction_task",
    } & node_ids
    assert ("finalize_qa", "finalize_delivery") in edges


def test_current_report_hierarchy_omits_reader_facing_evidence_notes(phase3_bundle):
    result = copy.deepcopy(phase3_bundle["result"])
    quality_profile = next(
        item
        for item in result["assessment_profile"]["dimensions"]
        if item["dimension"] == "Food Quality"
    )
    quality_profile["is_priority"] = True
    if "Food Quality" not in result["assessment_profile"][
        "priority_dimension_names"
    ]:
        result["assessment_profile"]["priority_dimension_names"].append(
            "Food Quality"
        )
    result["dimension_narratives"]["Food Quality"] = (
        fallback_dimension_narrative(
            quality_profile,
            assessment_profile=result["assessment_profile"],
        )
    )
    blocks = build_mfi_report_blocks(result)
    heading_texts = [
        block.text for block in blocks if block.type == "heading"
    ]
    for dimension in (
        "Assortment",
        "Availability",
        "Price",
        "Resilience",
        "Competition",
        "Infrastructure",
        "Service",
        "Food Quality",
        "Access & Protection",
    ):
        assert heading_texts.count(dimension) == 1
    assert not [block for block in blocks if block.type == "evidence_note"]
    cited_claim_count = 0
    for block in blocks:
        if (
            block.type == "paragraph"
            and block.meta
            and block.meta.get("claim_id")
            and (block.meta.get("metric_ids") or block.meta.get("document_ids"))
        ):
            cited_claim_count += 1
            assert isinstance(block.meta.get("metric_ids"), list)
            assert isinstance(block.meta.get("document_ids"), list)
    assert cited_claim_count > 0
    assert not [
        block
        for block in blocks
        if block.type == "figure" and block.figure_id == "risk_distribution"
    ]
    quality_tables = [
        block.meta.get("title")
        for block in blocks
        if block.type == "table"
        and block.meta
        and str(block.meta.get("title", "")).startswith("Food Quality:")
    ]
    assert "Food Quality: Ranked explanatory evidence" in quality_tables
    assert "Food Quality: Official subsection evidence" not in quality_tables
    docx = build_docx_bytes_from_report_blocks(blocks, visualizations={})
    document = Document(io.BytesIO(docx))
    document_text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "Assessment metadata and coverage" in document_text
    assert "Executive summary" in document_text
    assert "Methodology, limitations, and QA notices" in document_text
    assert "Evidence:" not in document_text


def test_historical_evidence_notes_remain_renderable(phase3_bundle):
    result = copy.deepcopy(phase3_bundle["result"])
    legacy_blocks = build_mfi_report_blocks(result, include_evidence_notes=True)
    evidence_notes = [
        block for block in legacy_blocks if block.type == "evidence_note"
    ]
    assert evidence_notes

    persisted = resolve_mfi_report_blocks(
        {"report_blocks": [block.model_dump() for block in legacy_blocks]}
    )
    assert persisted == legacy_blocks

    legacy_without_persisted_blocks = resolve_mfi_report_blocks(result)
    assert any(
        block.type == "evidence_note" for block in legacy_without_persisted_blocks
    )

    historical_docx = build_docx_bytes_from_report_blocks(
        [
            ReportBlock(
                type="evidence_note",
                text="Claim scope: assessed-market profile; Price mean: 5.00/10.",
            )
        ],
        visualizations={},
    )
    document = Document(io.BytesIO(historical_docx))
    assert any(
        paragraph.text.startswith("Evidence:") for paragraph in document.paragraphs
    )


def test_report_delivers_prominent_unresolved_qa_warning(phase3_bundle):
    result = copy.deepcopy(phase3_bundle["result"])
    result["qa_review"] = {
        "status": "completed_with_warnings",
        "correction_attempts": 3,
        "flags": [
            {
                "flag_id": "qa-1",
                "severity": "high",
                "message": "Unresolved claim.",
            }
        ],
    }
    blocks = build_mfi_report_blocks(result)
    warning = next(block for block in blocks if block.type == "qa_warning")
    assert "unresolved material issues" in warning.text


def test_graph_visuals_use_neutral_market_distribution(monkeypatch, phase3_bundle):
    def close_plot_and_return_image():
        import matplotlib.pyplot as plt

        plt.close()
        return "image"

    monkeypatch.setattr(graph, "save_plot_to_base64", close_plot_and_return_image)
    update = graph.node_mfi_graph_designer(
        {
            "country": "Testland",
            "assessment_profile": phase3_bundle["profile"],
            "markets_data": phase3_bundle["data"]["markets_data"],
        }
    )
    visualizations = update["visualizations"]
    assert visualizations["market_score_ranking"] == "image"
    assert "risk_distribution" not in visualizations
    assert visualizations["mfi_radar"] == "image"


def test_api_serialization_exposes_canonical_and_output_only_aliases(
    phase3_bundle,
):
    output = _build_mfi_output(
        result=phase3_bundle["result"],
        country="Testland",
        data_collection_start="2026-01-01",
        data_collection_end="2026-01-31",
    )
    assert output.narrative_schema_version == "2.0"
    assert output.release_control.analysis_version == "2"
    assert output.generation_diagnostics.executive_summary_mode == "fallback"
    assert len(output.dimension_narratives) == 9
    assert output.executive_summary_narrative.key_findings
    assert output.market_score_distribution
    assert output.dimension_findings
    assert all("sub_scores" in market for market in output.markets_data)
    assert all(
        "sub_scores" not in market
        for market in phase3_bundle["result"]["markets_data"]
    )
    dispatcher_output = _build_mfi_report_output(
        result=phase3_bundle["result"],
        run_id="phase3-test",
        country="Testland",
        data_collection_start="2026-01-01",
        data_collection_end="2026-01-31",
    )
    assert dispatcher_output["narrative_schema_version"] == "2.0"
    assert dispatcher_output["release_control"]["analysis_version"] == "2"
    assert dispatcher_output["generation_diagnostics"][
        "executive_summary_mode"
    ] == "fallback"
    assert len(dispatcher_output["dimension_narratives"]) == 9
    assert dispatcher_output["claim_validation"]["status"] == "passed"
    assert dispatcher_output["market_score_distribution"]


def test_report_graph_and_ui_have_no_internal_legacy_reads():
    root = Path(__file__).resolve().parents[1]
    paths = [
        root / "app/services/mfi_drafter/graph.py",
        root / "app/shared/report_blocks.py",
        root / "pages/4_MFI_Drafter.py",
    ]
    forbidden = ('"sub_scores"', '"risk_distribution"', '"risk_level"')
    for path in paths:
        text = path.read_text(encoding="utf-8")
        for token in forbidden:
            assert token not in text, f"{token} is read internally by {path.name}"
