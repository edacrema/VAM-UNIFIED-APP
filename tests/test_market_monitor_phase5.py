import io

from docx import Document

from app.shared.docx_export import build_docx_bytes_from_report_blocks
from app.shared.report_blocks import build_market_monitor_report_blocks
import streamlit_shared as shared


def _basket(role, name, description, scope="national", regions=None, version=1):
    return {
        "basket_role": role,
        "basket_version_id": f"{role}-v{version}",
        "version_number": version,
        "basket_name": name,
        "short_description": description,
        "scope_type": scope,
        "regions": regions or [],
        "items": [
            {
                "commodity_id": 1,
                "commodity_name_snapshot": "Maize",
                "databridges_unit_id": 100,
                "databridges_unit": "kg",
                "weight_quantity": 2,
                "sort_order": 1,
            },
            {
                "commodity_id": 2,
                "commodity_name_snapshot": "Beans",
                "databridges_unit_id": 100,
                "databridges_unit": "kg",
                "weight_quantity": 3.5,
                "sort_order": 2,
            },
        ],
    }


def _two_basket_result():
    primary = _basket("primary", "MEB", "Primary reference basket.")
    secondary = _basket(
        "secondary",
        "Pastoral basket",
        "Country Office pastoral livelihood basket.",
        scope="selected_regions",
        regions=["Region B", "Region A"],
        version=2,
    )
    return {
        "country": "South Sudan",
        "time_period": "2026-02",
        "language": "en",
        "secondary_basket_included": True,
        "food_basket": primary,
        "food_baskets": {"primary": primary, "secondary": secondary},
        "report_draft_sections": {
            "HIGHLIGHTS": "Highlights text.",
            "MARKET_OVERVIEW": "Overview text.",
            "COMMODITY_ANALYSIS": "Commodity text.",
            "REGIONAL_HIGHLIGHTS": "Regional text.\n\n[INSERT GRAPH: regional_comparison]",
        },
        "visualizations": {
            "food_basket_trend": "primary-trend",
            "food_basket_trend_primary": "primary-trend",
            "food_basket_trend_secondary": "secondary-trend",
            "regional_comparison": "primary-regional",
            "regional_comparison_primary": "primary-regional",
            "regional_comparison_secondary": "secondary-regional",
        },
    }


def test_report_assembly_orders_definition_table_and_available_role_figures():
    blocks = build_market_monitor_report_blocks(_two_basket_result())
    headings = [(index, block.text) for index, block in enumerate(blocks) if block.type == "heading"]
    table_index = next(index for index, block in enumerate(blocks) if block.type == "table")
    table = blocks[table_index]
    figures = [(index, block.figure_id) for index, block in enumerate(blocks) if block.type == "figure"]

    highlight_index = next(index for index, text in headings if text == "Highlights")
    definition_index = next(index for index, text in headings if text == "Basket Definitions")
    overview_index = next(index for index, text in headings if text == "Market Overview")
    primary_trend_index = next(index for index, figure in figures if figure == "food_basket_trend_primary")
    secondary_trend_index = next(index for index, figure in figures if figure == "food_basket_trend_secondary")
    assert highlight_index < definition_index < table_index < primary_trend_index < secondary_trend_index < overview_index
    assert [figure for _index, figure in figures].count("regional_comparison_primary") == 1
    assert [figure for _index, figure in figures].count("regional_comparison_secondary") == 1
    assert "food_basket_trend" not in [figure for _index, figure in figures]

    assert table.meta["table_kind"] == "basket_definitions"
    assert [row["basket_role"] for row in table.meta["rows"]] == ["primary", "secondary"]
    secondary = table.meta["rows"][1]
    assert secondary["regions"] == ["Region B", "Region A"]
    assert secondary["basket_version_id"] == "secondary-v2"
    assert [item["quantity"] for item in secondary["components"]] == [2.0, 3.5]
    assert secondary["display"]["composition"] == "2 kg Maize; 3.5 kg Beans"


def test_primary_only_legacy_result_gets_definition_and_legacy_figure_fallback():
    primary = _basket("primary", "Custom basket", "CO-authored description.")
    blocks = build_market_monitor_report_blocks(
        {
            "country": "Testland",
            "time_period": "2026-02",
            "food_basket": primary,
            "food_baskets": {},
            "report_sections": {"HIGHLIGHTS": "Text."},
            "visualizations": {"food_basket_trend": "legacy"},
        }
    )

    table = next(block for block in blocks if block.type == "table")
    figures = [block.figure_id for block in blocks if block.type == "figure"]
    assert len(table.meta["rows"]) == 1
    assert table.meta["rows"][0]["basket_name"] == "Custom basket"
    assert figures == ["food_basket_trend"]


def test_report_assembly_suppresses_unproduced_inline_and_secondary_figures():
    primary = _basket("primary", "MEB", "Primary reference basket.")
    blocks = build_market_monitor_report_blocks(
        {
            "country": "Testland",
            "time_period": "2026-02",
            "food_baskets": {"primary": primary, "secondary": None},
            "secondary_basket_included": False,
            "report_sections": {
                "REGIONAL_HIGHLIGHTS": "Text. [INSERT GRAPH: regional_comparison]",
            },
            "visualizations": {},
        }
    )

    assert not [block for block in blocks if block.type == "figure"]


def test_basket_table_localizes_labels_without_translating_user_content():
    result = _two_basket_result()
    result["language"] = "fr"
    table = next(block for block in build_market_monitor_report_blocks(result) if block.type == "table")

    assert table.meta["headers"]["basket"] == "Panier / role"
    assert table.meta["rows"][1]["basket_name"] == "Pastoral basket"
    assert table.meta["rows"][1]["short_description"] == "Country Office pastoral livelihood basket."
    assert table.meta["rows"][1]["scope_label"].startswith("Regions selectionnees")


def test_docx_renders_basket_definition_table_without_missing_figure_placeholders():
    result = _two_basket_result()
    result["visualizations"] = {}
    result["report_draft_sections"] = {}
    blocks = build_market_monitor_report_blocks(result)
    document = Document(io.BytesIO(build_docx_bytes_from_report_blocks(blocks, visualizations={})))

    basket_table = next(table for table in document.tables if table.cell(0, 0).text == "Basket / role")
    assert len(basket_table.rows) == 3
    assert basket_table.cell(1, 0).text == "MEB (Primary)"
    assert basket_table.cell(2, 0).text == "Pastoral basket (Secondary)"
    assert "3.5 kg Beans" in basket_table.cell(2, 3).text
    assert len(document.inline_shapes) == 0


def test_streamlit_renderers_show_read_only_basket_table_and_deduplicate_aliases(monkeypatch):
    result = _two_basket_result()
    table = next(block for block in build_market_monitor_report_blocks(result) if block.type == "table")
    rendered_frames = []
    rendered_images = []
    rendered_headers = []

    monkeypatch.setattr(shared.st, "dataframe", lambda frame, **_kwargs: rendered_frames.append(frame))
    monkeypatch.setattr(shared.st, "write", lambda *_args, **_kwargs: None)
    monkeypatch.setattr(shared.st, "subheader", lambda text: rendered_headers.append(text))
    monkeypatch.setattr(shared.st, "image", lambda image, **kwargs: rendered_images.append((image, kwargs)))
    monkeypatch.setattr(shared, "decode_base64_data", lambda value: str(value).encode("utf-8"))

    shared.render_report_blocks([table.model_dump(), {"type": "figure", "figure_id": "missing"}], {})
    shared.render_visualizations(result["visualizations"])

    assert rendered_frames[0].columns.tolist() == ["Basket / role", "Description", "Scope", "Composition"]
    assert rendered_frames[0].iloc[1, 0] == "Pastoral basket (Secondary)"
    assert "food_basket_trend" not in rendered_headers
    assert "regional_comparison" not in rendered_headers
    assert "food_basket_trend_primary" in rendered_headers
    assert "regional_comparison_primary" in rendered_headers
    assert len(rendered_images) == 4


def test_dispatcher_service_info_documents_phase5_visualization_contract():
    response = shared.dispatch_request("GET", "/market-monitor/info")

    assert response.status_code == 200
    assert response.json()["basket_visualizations"]["canonical"] == [
        "food_basket_trend_primary",
        "food_basket_trend_secondary",
        "regional_comparison_primary",
        "regional_comparison_secondary",
    ]
    assert response.json()["basket_visualizations"]["combined_chart"] is False
    assert "qa_review" in response.json()["outputs"]
