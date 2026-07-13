"""Phase 7 release qualification for the Market Monitor second basket.

The migration commands intentionally operate only on a copied SQLite file or
an operator-provided disposable PostgreSQL clone. The sample command writes
synthetic, non-publication QA reports under an ignored output directory.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import shutil
import sys
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Mapping, Sequence

import pandas as pd
from docx import Document
from sqlalchemy import create_engine, inspect, text
from sqlalchemy.engine import Engine, make_url

REPO_ROOT = Path(__file__).resolve().parents[1]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from app.services.market_monitor.basket_calculation import (
    BasketCalculationSpec,
    calculate_basket_series,
    calculate_basket_statistics,
)
from app.services.market_monitor.food_basket import (
    BasketRole,
    BasketScopeType,
    DEFAULT_PRIMARY_BASKET_DESCRIPTION,
    SqlCountryFoodBasketRepository,
)
from app.services.market_monitor.graph import node_graph_designer
from app.services.price_cache.config import load_price_cache_config
from app.services.price_cache.fixtures import seed_cache_snapshot
from app.services.price_cache.migrations import apply_migrations
from app.services.price_cache.sql_repository import SqlPriceCacheRepository, create_price_cache_engine
from app.shared.docx_export import build_docx_bytes_from_report_blocks
from app.shared.report_blocks import build_market_monitor_report_blocks


MIGRATION_004 = "004_country_food_baskets"
MIGRATION_005 = "005_second_food_basket"
DISPOSABLE_CONFIRMATION = "DISPOSABLE"
DEFAULT_OUTPUT_DIR = Path(".tmp/phase7-second-basket")


class ReleaseQAFailure(RuntimeError):
    """Raised when a release qualification invariant is not satisfied."""


def verify_sqlite_copy(source: Path, output_dir: Path) -> dict[str, Any]:
    """Copy and migrate a populated v004 SQLite database without touching it."""

    source = source.expanduser().resolve()
    if not source.is_file():
        raise ReleaseQAFailure(f"SQLite source does not exist: {source}")
    source_engine = create_engine(f"sqlite:///{source.as_posix()}", future=True)
    try:
        _preflight_v004(source_engine)
    finally:
        source_engine.dispose()

    output_dir.mkdir(parents=True, exist_ok=True)
    stamp = _utc_stamp()
    target = (output_dir / f"sqlite-v004-rehearsal-{stamp}.sqlite3").resolve()
    if target == source:
        raise ReleaseQAFailure("The SQLite rehearsal target must differ from the source database.")
    shutil.copy2(source, target)

    engine = create_engine(f"sqlite:///{target.as_posix()}", future=True)
    try:
        audit = verify_populated_v004_migration(engine, "sqlite")
    finally:
        engine.dispose()
    audit.update({"backend": "sqlite", "target_copy": target.name})
    _write_json(output_dir / f"sqlite-migration-audit-{stamp}.json", audit)
    return audit


def verify_postgres_clone(database_url: str, output_dir: Path, confirmation: str) -> dict[str, Any]:
    """Migrate an operator-restored disposable PostgreSQL clone in place."""

    if confirmation != DISPOSABLE_CONFIRMATION:
        raise ReleaseQAFailure(
            f"PostgreSQL rehearsal requires --confirm-disposable {DISPOSABLE_CONFIRMATION}."
        )
    if not str(database_url or "").strip():
        raise ReleaseQAFailure("A disposable PostgreSQL clone URL is required.")
    configured_url = str(os.getenv("PRICE_CACHE_DATABASE_URL") or "").strip()
    if configured_url and _safe_url(configured_url) == _safe_url(database_url):
        raise ReleaseQAFailure(
            "The rehearsal URL matches PRICE_CACHE_DATABASE_URL; restore and use a separate disposable clone."
        )

    engine = create_engine(database_url, future=True)
    try:
        audit = verify_populated_v004_migration(engine, "postgres")
    finally:
        engine.dispose()
    audit.update({"backend": "postgres", "target": "operator_confirmed_disposable_clone"})
    output_dir.mkdir(parents=True, exist_ok=True)
    _write_json(output_dir / f"postgres-migration-audit-{_utc_stamp()}.json", audit)
    return audit


def verify_populated_v004_migration(engine: Engine, dialect: str) -> dict[str, Any]:
    """Verify ID-preserving migration and role-aware repository behavior."""

    before = _preflight_v004(engine)
    applied = apply_migrations(engine, dialect)
    applied_again = apply_migrations(engine, dialect)
    if applied != [MIGRATION_005]:
        raise ReleaseQAFailure(f"Expected only {MIGRATION_005}; applied {applied!r}.")
    if applied_again:
        raise ReleaseQAFailure(f"Second migration run was not a no-op: {applied_again!r}.")

    after = _snapshot_legacy_rows(engine)
    for key in (
        "version_count",
        "item_count",
        "pointer_count",
        "version_id_hash",
        "item_id_hash",
        "pointer_id_hash",
    ):
        if after[key] != before[key]:
            raise ReleaseQAFailure(f"Legacy preservation failed for {key}.")

    with engine.begin() as conn:
        invalid_versions = conn.execute(
            text(
                """
                SELECT COUNT(*)
                FROM country_food_basket_versions
                WHERE basket_role <> 'primary'
                   OR basket_name <> 'MEB'
                   OR scope_type <> 'national'
                   OR short_description IS NULL
                """
            )
        ).scalar_one()
        invalid_pointers = conn.execute(
            text("SELECT COUNT(*) FROM country_food_basket_current WHERE basket_role <> 'primary'")
        ).scalar_one()
    if int(invalid_versions or 0) or int(invalid_pointers or 0):
        raise ReleaseQAFailure("Legacy basket metadata or pointers were not backfilled as primary MEB records.")

    schema = _schema_audit(engine)
    if not schema["valid"]:
        raise ReleaseQAFailure("Role-aware schema constraints or indexes are incomplete.")

    repository = SqlCountryFoodBasketRepository(engine)
    price_repository = SqlPriceCacheRepository(engine)
    country = before["sample_country"]
    primary = repository.get_active_basket(country, BasketRole.PRIMARY)
    if primary is None:
        raise ReleaseQAFailure(f"No active legacy primary basket is available for {country}.")
    availability = price_repository.get_country_availability(country)
    priced_ids = list(getattr(availability, "priced_commodity_ids", None) or [])
    if not priced_ids:
        raise ReleaseQAFailure(f"No active priced commodity is available for repository smoke testing in {country}.")
    available_regions = list(getattr(availability, "admin1_names", None) or [])
    if not available_regions:
        raise ReleaseQAFailure(f"No active Admin1 region is available for region-storage smoke testing in {country}.")

    secondary = repository.save_basket(
        country,
        role=BasketRole.SECONDARY,
        basket_name="Phase 7 rehearsal basket",
        short_description="Disposable migration rehearsal basket.",
        scope_type=BasketScopeType.SELECTED_REGIONS,
        regions=[available_regions[0]],
        items=[{"commodity_id": int(priced_ids[0]), "weight_quantity": 1}],
        created_by_user_id="phase7-release-qa",
        change_note="Disposable role-aware repository smoke test.",
    )
    active = repository.get_active_baskets(country)
    if active["primary"] is None or active["primary"].basket_version_id != primary.basket_version_id:
        raise ReleaseQAFailure("Saving a secondary changed the active primary pointer.")
    if active["secondary"] is None or active["secondary"].basket_version_id != secondary.basket_version_id:
        raise ReleaseQAFailure("The secondary pointer was not created.")
    if secondary.version_number != int(before["max_version_number"]) + 1:
        raise ReleaseQAFailure("Country-wide version numbering did not advance across roles.")
    archived = repository.archive_secondary_basket(country)
    if archived is None or repository.archive_secondary_basket(country) is not None:
        raise ReleaseQAFailure("Secondary archive is not role-specific and idempotent.")
    historical = repository.get_basket_version(country, secondary.basket_version_id)
    if historical is None or historical.status != "archived":
        raise ReleaseQAFailure("Archived secondary history is not retrievable by immutable ID.")
    if historical.regions != [available_regions[0]]:
        raise ReleaseQAFailure("Ordered canonical basket regions were not retained in history.")

    return {
        "status": "passed",
        "source_schema": MIGRATION_004,
        "applied": applied,
        "second_run_applied": applied_again,
        "legacy": {
            "version_count": before["version_count"],
            "item_count": before["item_count"],
            "pointer_count": before["pointer_count"],
            "version_id_hash": before["version_id_hash"],
            "item_id_hash": before["item_id_hash"],
            "pointer_id_hash": before["pointer_id_hash"],
            "identifiers_preserved": True,
        },
        "schema": schema,
        "repository_smoke": {
            "primary_preserved": True,
            "secondary_created": True,
            "country_wide_sequence": True,
            "archive_idempotent": True,
            "history_retrievable": True,
            "ordered_region_storage": True,
        },
        "completed_at": datetime.now(timezone.utc).isoformat(),
    }


def generate_sample_reports(output_dir: Path) -> dict[str, Any]:
    """Generate the three deterministic multilingual visual-QA reports."""

    output_dir.mkdir(parents=True, exist_ok=True)
    scenarios = [
        ("primary-only-en", "en", "primary_only"),
        ("two-national-fr", "fr", "two_national"),
        ("national-regional-es", "es", "national_regional"),
    ]
    results: list[dict[str, Any]] = []
    for slug, language, scenario_kind in scenarios:
        result = _build_sample_result(language=language, scenario_kind=scenario_kind)
        blocks = build_market_monitor_report_blocks(result)
        result["report_blocks"] = [block.model_dump() for block in blocks]
        docx_bytes = build_docx_bytes_from_report_blocks(
            blocks,
            visualizations=result["visualizations"],
            language=language,
        )

        docx_path = output_dir / f"{slug}.docx"
        json_path = output_dir / f"{slug}.json"
        docx_path.write_bytes(docx_bytes)
        _write_json(json_path, result)
        structural = _inspect_sample_docx(docx_path, result)
        if not structural["passed"]:
            raise ReleaseQAFailure(f"Structural DOCX QA failed for {slug}: {structural['checks']!r}")
        results.append(
            {
                "slug": slug,
                "language": language,
                "scenario": scenario_kind,
                "docx": docx_path.name,
                "json": json_path.name,
                "structural": structural,
                "visual_review": {"status": "pending", "reviewer": None, "reviewed_at": None},
            }
        )

    manifest = {
        "status": "pending_visual_review",
        "synthetic_data": True,
        "not_for_publication": True,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "scenarios": results,
    }
    _write_json(output_dir / "sample-report-manifest.json", manifest)
    (output_dir / "visual-review-checklist.md").write_text(
        _visual_checklist_markdown(manifest), encoding="utf-8"
    )
    return manifest


def prepare_streamlit_fixture(output_dir: Path) -> dict[str, Any]:
    """Create a deterministic, non-sensitive two-basket database for UI evidence."""

    output_dir.mkdir(parents=True, exist_ok=True)
    database_path = (output_dir / "streamlit-two-basket.sqlite3").resolve()
    database_path.unlink(missing_ok=True)
    config = load_price_cache_config(
        {"PRICE_CACHE_BACKEND": "sqlite", "PRICE_CACHE_SQLITE_PATH": str(database_path)}
    )
    engine = create_price_cache_engine(config)
    try:
        apply_migrations(engine, "sqlite")
        seed_cache_snapshot(SqlPriceCacheRepository(engine))
        repository = SqlCountryFoodBasketRepository(engine)
        primary = repository.save_basket(
            "SSD",
            role=BasketRole.PRIMARY,
            basket_name="MEB",
            short_description=DEFAULT_PRIMARY_BASKET_DESCRIPTION,
            items=[
                {"commodity_id": 1, "weight_quantity": 12, "item_note": "Cereal staple"},
                {"commodity_id": 2, "weight_quantity": 6, "item_note": "Pulse component"},
            ],
            created_by_user_id="phase7-synthetic-qa",
            change_note="Synthetic documentation example.",
        )
        secondary = repository.save_basket(
            "SSD",
            role=BasketRole.SECONDARY,
            basket_name="Urban reference basket",
            short_description="Synthetic affordability basket for selected urban monitoring regions.",
            scope_type=BasketScopeType.SELECTED_REGIONS,
            regions=["Central Equatoria"],
            items=[
                {"commodity_id": 1, "weight_quantity": 8, "item_note": "Urban cereal component"},
                {"commodity_id": 2, "weight_quantity": 4, "item_note": "Urban pulse component"},
            ],
            created_by_user_id="phase7-synthetic-qa",
            change_note="Synthetic documentation example.",
        )
    finally:
        engine.dispose()
    summary = {
        "status": "ready",
        "synthetic_data": True,
        "not_for_publication": True,
        "database": database_path.name,
        "country": "South Sudan",
        "primary_basket_version_id": primary.basket_version_id,
        "secondary_basket_version_id": secondary.basket_version_id,
    }
    _write_json(output_dir / "streamlit-fixture.json", summary)
    return summary


def record_visual_review(output_dir: Path, reviewer: str, status: str, notes: str) -> dict[str, Any]:
    """Record the human/render review after every generated page was inspected."""

    manifest_path = output_dir / "sample-report-manifest.json"
    if not manifest_path.is_file():
        raise ReleaseQAFailure(f"Sample manifest not found: {manifest_path}")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    reviewed_at = datetime.now(timezone.utc).isoformat()
    for scenario in manifest.get("scenarios") or []:
        scenario["visual_review"] = {
            "status": status,
            "reviewer": reviewer,
            "reviewed_at": reviewed_at,
            "notes": notes,
        }
    manifest["status"] = "passed" if status == "passed" else "failed"
    manifest["reviewed_at"] = reviewed_at
    _write_json(manifest_path, manifest)
    (output_dir / "visual-review-checklist.md").write_text(
        _visual_checklist_markdown(manifest), encoding="utf-8"
    )
    return manifest


def _preflight_v004(engine: Engine) -> dict[str, Any]:
    inspector = inspect(engine)
    required = {
        "price_cache_schema_migrations",
        "country_food_basket_versions",
        "country_food_basket_items",
        "country_food_basket_current",
    }
    missing = sorted(required - set(inspector.get_table_names()))
    if missing:
        raise ReleaseQAFailure("Migration source is missing required tables: " + ", ".join(missing))
    with engine.begin() as conn:
        migrations = {
            str(row[0])
            for row in conn.execute(text("SELECT version FROM price_cache_schema_migrations")).all()
        }
    if MIGRATION_004 not in migrations or MIGRATION_005 in migrations:
        raise ReleaseQAFailure("Migration source must be populated at schema version 004 and not yet contain 005.")
    snapshot = _snapshot_legacy_rows(engine)
    if min(snapshot["version_count"], snapshot["item_count"], snapshot["pointer_count"]) <= 0:
        raise ReleaseQAFailure("Migration source must contain legacy basket versions, items, and current pointers.")
    if snapshot["sample_country"] is None:
        raise ReleaseQAFailure("No country has both a legacy version and current pointer.")
    return snapshot


def _snapshot_legacy_rows(engine: Engine) -> dict[str, Any]:
    with engine.begin() as conn:
        version_rows = conn.execute(
            text("SELECT basket_version_id, country_iso3, version_number FROM country_food_basket_versions")
        ).all()
        item_rows = conn.execute(text("SELECT basket_item_id FROM country_food_basket_items")).all()
        pointer_rows = conn.execute(
            text("SELECT country_iso3, active_basket_version_id FROM country_food_basket_current")
        ).all()
    versions = sorted(str(row[0]) for row in version_rows)
    items = sorted(str(row[0]) for row in item_rows)
    pointers = sorted(f"{row[0]}|{row[1]}" for row in pointer_rows)
    countries = sorted({str(row[1]).upper() for row in version_rows})
    pointer_countries = sorted({str(row[0]).upper() for row in pointer_rows})
    sample_country = next((country for country in pointer_countries if country in countries), None)
    max_version_number = (
        max(int(row[2]) for row in version_rows if str(row[1]).upper() == sample_country)
        if sample_country is not None
        else 0
    )
    return {
        "version_count": len(version_rows),
        "item_count": len(item_rows),
        "pointer_count": len(pointer_rows),
        "version_id_hash": _identifier_hash(versions),
        "item_id_hash": _identifier_hash(items),
        "pointer_id_hash": _identifier_hash(pointers),
        "max_version_number": max_version_number,
        "sample_country": sample_country,
    }


def _schema_audit(engine: Engine) -> dict[str, Any]:
    inspector = inspect(engine)
    tables = set(inspector.get_table_names())
    pk_columns = inspector.get_pk_constraint("country_food_basket_current").get("constrained_columns") or []
    foreign_keys = inspector.get_foreign_keys("country_food_basket_current")
    role_fk = any(
        list(item.get("constrained_columns") or [])
        == ["country_iso3", "basket_role", "active_basket_version_id"]
        and list(item.get("referred_columns") or [])
        == ["country_iso3", "basket_role", "basket_version_id"]
        for item in foreign_keys
    )
    version_indexes = {item["name"] for item in inspector.get_indexes("country_food_basket_versions")}
    region_indexes = (
        {item["name"] for item in inspector.get_indexes("country_food_basket_regions")}
        if "country_food_basket_regions" in tables
        else set()
    )
    checks = {
        "region_table": "country_food_basket_regions" in tables,
        "composite_pointer_primary_key": pk_columns == ["country_iso3", "basket_role"],
        "role_matched_pointer_foreign_key": role_fk,
        "role_history_index": "idx_country_food_basket_versions_role_history" in version_indexes,
        "ordered_region_index": "idx_country_food_basket_regions_version" in region_indexes,
    }
    return {"valid": all(checks.values()), "checks": checks}


def _build_sample_result(*, language: str, scenario_kind: str) -> dict[str, Any]:
    months = pd.date_range("2025-01-01", periods=13, freq="MS")
    target = months[-1]
    available_regions = ["Region A", "Region B"]
    price_frame = _synthetic_price_frame(months)
    primary = _basket_snapshot(
        role="primary",
        version="primary-v1",
        name="MEB",
        description=DEFAULT_PRIMARY_BASKET_DESCRIPTION,
        scope="national",
        regions=[],
        items=[(1, "Maize", 2.0), (2, "Beans", 3.0)],
    )
    secondary: dict[str, Any] | None = None
    if scenario_kind == "two_national":
        secondary = _basket_snapshot(
            role="secondary",
            version="secondary-v2",
            name="Panier pastoral",
            description="Panier de référence défini par le Bureau pays pour les ménages pastoraux.",
            scope="national",
            regions=[],
            items=[(1, "Maize", 1.0), (3, "Vegetable oil", 4.0)],
        )
    elif scenario_kind == "national_regional":
        secondary = _basket_snapshot(
            role="secondary",
            version="secondary-v2",
            name="Canasta urbana regional",
            description="Canasta definida por la Oficina de País para el seguimiento urbano regional.",
            scope="selected_regions",
            regions=available_regions,
            items=[(2, "Beans", 2.0), (3, "Vegetable oil", 2.5)],
        )

    snapshots = [primary] + ([secondary] if secondary else [])
    specs = [BasketCalculationSpec.from_snapshot(item) for item in snapshots]
    calculation = calculate_basket_series(
        price_frame,
        specs,
        full_date_index=pd.DatetimeIndex(months),
        run_regions=available_regions,
        available_regions=available_regions,
    )
    statistics = calculate_basket_statistics(price_frame, specs, calculation, target_date=target)
    national_prices = (
        price_frame.groupby(["Price Date", "Commodity"], as_index=False)["Price"].mean()
        .pivot(index="Price Date", columns="Commodity", values="Price")
        .sort_index()
    )
    regional_prices = calculation.regional[calculation.regional["BasketRole"] == "primary"]
    regional_prices = regional_prices[["Date", "Region", "Cost"]].rename(columns={"Cost": "FoodBasket"})
    state = {
        "country": "South Sudan",
        "time_period": target.strftime("%Y-%m"),
        "language": language,
        "currency_code": "SSP",
        "time_series_data_national": national_prices.to_json(date_format="iso"),
        "time_series_history_national": "",
        "time_series_data_regional": regional_prices.to_json(orient="records", date_format="iso"),
        "food_basket": primary,
        "food_baskets": {"primary": primary, "secondary": secondary},
        "include_secondary_basket": secondary is not None,
        "secondary_basket_included": secondary is not None,
        "basket_statistics": statistics,
        "basket_series_national": calculation.records("national"),
        "basket_series_regional": calculation.records("regional"),
        "data_statistics": {"food_basket": statistics["primary"] or {}, "commodities": {}},
    }
    visualizations = node_graph_designer(state)["visualizations"]
    primary_alias = dict(statistics["primary"] or {})
    primary_alias["current_price"] = primary_alias.get("current_cost")
    sections = _sample_sections(language, secondary)
    return {
        "run_id": f"phase7-{scenario_kind}-{language}",
        "country": "South Sudan",
        "time_period": target.strftime("%Y-%m"),
        "language": language,
        "locale": {"en": "en_US", "fr": "fr_FR", "es": "es_ES"}[language],
        "report_draft_sections": sections,
        "report_sections": sections,
        "visualizations": visualizations,
        "data_statistics": {"food_basket": primary_alias, "commodities": {}},
        "food_basket": primary,
        "food_baskets": {"primary": primary, "secondary": secondary},
        "basket_statistics": statistics,
        "basket_series_national": calculation.records("national"),
        "basket_series_regional": calculation.records("regional"),
        "secondary_basket_included": secondary is not None,
        "qa_review": {"status": "passed", "correction_attempts": 0, "flags": []},
        "warnings": calculation.warnings,
        "document_references": [],
        "module_sections": {},
        "success": True,
        "synthetic_qa_fixture": True,
    }


def _synthetic_price_frame(months: Sequence[pd.Timestamp]) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    commodities = [(1, "Maize", 12.0), (2, "Beans", 8.0), (3, "Vegetable oil", 4.0)]
    for month_index, month in enumerate(months):
        for region_index, region in enumerate(("Region A", "Region B")):
            for commodity_id, commodity_name, base in commodities:
                if month_index == 5 and commodity_id == 2:
                    continue
                if month_index == 8 and region == "Region B" and commodity_id == 3:
                    continue
                rows.append(
                    {
                        "Commodity ID": commodity_id,
                        "Commodity": commodity_name,
                        "Price Date": pd.Timestamp(month),
                        "Price": base + month_index * 0.35 + region_index * 1.5,
                        "Admin 1": region,
                        "Unit ID": 100,
                        "Unit": "kg",
                    }
                )
    return pd.DataFrame(rows)


def _basket_snapshot(
    *,
    role: str,
    version: str,
    name: str,
    description: str,
    scope: str,
    regions: Sequence[str],
    items: Sequence[tuple[int, str, float]],
) -> dict[str, Any]:
    return {
        "basket_role": role,
        "basket_version_id": version,
        "version_number": 1 if role == "primary" else 2,
        "basket_name": name,
        "short_description": description,
        "scope_type": scope,
        "regions": list(regions),
        "status": "active",
        "created_at": "2026-07-01T00:00:00+00:00",
        "created_by_user_id": "phase7-release-qa",
        "cache_version_id_at_creation": "synthetic-cache-v1",
        "change_note": "Synthetic Phase 7 release qualification fixture.",
        "items": [
            {
                "basket_item_id": f"{version}-item-{commodity_id}",
                "basket_version_id": version,
                "commodity_id": commodity_id,
                "commodity_name_snapshot": commodity_name,
                "databridges_unit_id": 100,
                "databridges_unit": "kg",
                "weight_quantity": quantity,
                "sort_order": index,
                "item_note": None,
            }
            for index, (commodity_id, commodity_name, quantity) in enumerate(items, start=1)
        ],
    }


def _sample_sections(language: str, secondary: Mapping[str, Any] | None) -> dict[str, str]:
    secondary_name = str((secondary or {}).get("basket_name") or "")
    if language == "fr":
        return {
            "HIGHLIGHTS": "Le MEB et le Panier pastoral sont présentés séparément selon leur composition immuable.",
            "MARKET_OVERVIEW": (
                "Le MEB constitue le panier principal national. Le Panier pastoral constitue un second panier "
                "national défini par le Bureau pays; leurs coûts absolus ne sont pas comparés directement."
            ),
            "COMMODITY_ANALYSIS": "Les contributions reflètent la composition des coûts et non une causalité de mouvement.",
            "REGIONAL_HIGHLIGHTS": "Les résultats régionaux complets sont affichés pour chaque panier.\n\n[INSERT GRAPH: regional_comparison]",
        }
    if language == "es":
        return {
            "HIGHLIGHTS": f"El MEB y la {secondary_name} se presentan con identidades y alcances separados.",
            "MARKET_OVERVIEW": (
                "El MEB conserva su alcance nacional. La Canasta urbana regional se calcula solamente para "
                "Region A y Region B, sin presentarla como un valor nacional."
            ),
            "COMMODITY_ANALYSIS": "Las contribuciones describen la composición del costo y no prueban causalidad.",
            "REGIONAL_HIGHLIGHTS": "Los gráficos omiten observaciones incompletas.\n\n[INSERT GRAPH: regional_comparison]",
        }
    return {
        "HIGHLIGHTS": "The national MEB is reported using its immutable definition and complete-component costs.",
        "MARKET_OVERVIEW": "The MEB remains the required primary national basket for this synthetic QA report.",
        "COMMODITY_ANALYSIS": "Component contributions describe cost composition rather than movement causality.",
        "REGIONAL_HIGHLIGHTS": "Complete regional breakdowns are shown for the selected month.\n\n[INSERT GRAPH: regional_comparison]",
    }


def _inspect_sample_docx(path: Path, result: Mapping[str, Any]) -> dict[str, Any]:
    document = Document(path)
    expected_rows = 3 if result.get("secondary_basket_included") else 2
    basket_tables = [
        table
        for table in document.tables
        if table.rows and table.cell(0, 0).text in {"Basket / role", "Panier / role", "Canasta / rol"}
    ]
    canonical_figures = [
        key
        for key in (
            "food_basket_trend_primary",
            "food_basket_trend_secondary",
            "regional_comparison_primary",
            "regional_comparison_secondary",
        )
        if key in (result.get("visualizations") or {})
    ]
    checks = {
        "basket_definition_table": len(basket_tables) == 1,
        "basket_definition_rows": bool(basket_tables) and len(basket_tables[0].rows) == expected_rows,
        "canonical_figures_generated": len(canonical_figures) == (4 if result.get("secondary_basket_included") else 2),
        "figures_embedded": len(document.inline_shapes) >= len(canonical_figures),
        "co_authored_text_present": all(
            str(snapshot.get("short_description") or "") in "\n".join(
                cell.text for table in document.tables for row in table.rows for cell in row.cells
            )
            for snapshot in (result.get("food_baskets") or {}).values()
            if isinstance(snapshot, Mapping)
        ),
    }
    return {"passed": all(checks.values()), "checks": checks, "inline_shapes": len(document.inline_shapes)}


def _visual_checklist_markdown(manifest: Mapping[str, Any]) -> str:
    status = str(manifest.get("status") or "pending_visual_review")
    lines = [
        "# Phase 7 sample-report visual review",
        "",
        f"Overall status: **{status}**",
        "",
        "These reports use deterministic synthetic data and are not for publication.",
        "",
    ]
    for scenario in manifest.get("scenarios") or []:
        review = scenario.get("visual_review") or {}
        checked = "x" if review.get("status") == "passed" else " "
        lines.extend(
            [
                f"## {scenario.get('slug')}",
                "",
                f"- [{checked}] Every page inspected at 100% zoom",
                f"- [{checked}] Basket definition rows, descriptions, scope, and composition are readable",
                f"- [{checked}] Chart titles and captions identify the basket and scope",
                f"- [{checked}] Incomplete historical values appear as gaps, not partial costs",
                f"- [{checked}] No duplicate compatibility figures or empty placeholders",
                f"- [{checked}] Localized labels render correctly and CO-authored text is unchanged",
                f"- Status: {review.get('status') or 'pending'}",
                f"- Reviewer: {review.get('reviewer') or 'pending'}",
                f"- Reviewed at: {review.get('reviewed_at') or 'pending'}",
                f"- Notes: {review.get('notes') or ''}",
                "",
            ]
        )
    return "\n".join(lines)


def _identifier_hash(values: Sequence[str]) -> str:
    return hashlib.sha256("\n".join(values).encode("utf-8")).hexdigest()


def _safe_url(value: str) -> str:
    return make_url(value).render_as_string(hide_password=True)


def _utc_stamp() -> str:
    return datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")


def _write_json(path: Path, payload: Any) -> None:
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, default=_json_default),
        encoding="utf-8",
    )


def _json_default(value: Any) -> Any:
    if isinstance(value, pd.Timestamp):
        return value.isoformat()
    if hasattr(value, "item"):
        return value.item()
    if hasattr(value, "model_dump"):
        return value.model_dump()
    raise TypeError(f"Cannot JSON encode {type(value).__name__}")


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)

    sqlite_parser = subparsers.add_parser("verify-sqlite", help="Copy and migrate a populated SQLite v004 database.")
    sqlite_parser.add_argument("--source", type=Path, required=True)
    sqlite_parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)

    postgres_parser = subparsers.add_parser(
        "verify-postgres", help="Migrate an operator-restored disposable PostgreSQL v004 clone."
    )
    postgres_parser.add_argument(
        "--database-url",
        default=os.getenv("PHASE7_POSTGRES_CLONE_DATABASE_URL", ""),
    )
    postgres_parser.add_argument("--confirm-disposable", default="")
    postgres_parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)

    samples_parser = subparsers.add_parser("generate-samples", help="Generate three synthetic multilingual QA reports.")
    samples_parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)

    ui_parser = subparsers.add_parser(
        "prepare-ui-fixture",
        help="Create an ignored synthetic two-basket SQLite database for Streamlit screenshots.",
    )
    ui_parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)

    review_parser = subparsers.add_parser("record-review", help="Record the completed visual-review outcome.")
    review_parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    review_parser.add_argument("--reviewer", required=True)
    review_parser.add_argument("--status", choices=["passed", "failed"], required=True)
    review_parser.add_argument("--notes", default="")
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = _build_parser().parse_args(argv)
    if args.command == "verify-sqlite":
        result = verify_sqlite_copy(args.source, args.output_dir)
    elif args.command == "verify-postgres":
        result = verify_postgres_clone(
            args.database_url,
            args.output_dir,
            args.confirm_disposable,
        )
    elif args.command == "generate-samples":
        result = generate_sample_reports(args.output_dir)
    elif args.command == "prepare-ui-fixture":
        result = prepare_streamlit_fixture(args.output_dir)
    else:
        result = record_visual_review(args.output_dir, args.reviewer, args.status, args.notes)
    print(json.dumps(result, ensure_ascii=False, indent=2, default=_json_default))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
